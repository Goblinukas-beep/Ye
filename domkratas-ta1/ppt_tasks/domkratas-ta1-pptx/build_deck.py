"""Build the editable PPTX for 'Domkratas' (PBL TA1).

Reproducible Build Mode script (ppt-design-skill + pptx-designer public API).
Design direction: theme-lock.yaml v1 (graphite + safety yellow). Coordinates are
authored on the same 1920×1080 grid as the HTML deck and converted with px().
Schemes are PNGs rendered from ../../src/diagrams.py (node ../../src/render.js svg ...).

Run:  python build_deck.py   → ../../domkratas-ta1.pptx
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from PIL import Image, ImageDraw
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pptx_designer import Presentation, validate_resolved_theme
from pptx_designer.renderer.theme import ThemeComposer
from pptx_designer.tools.images import cover_image
from pptx_designer.tools.shapes import add_slide, oval, rect
from pptx_designer.tools.text import text

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUT = ROOT / "domkratas-ta1.pptx"
PNG = ROOT / "src" / "png"
ASSETS = HERE / "assets"

# ── Locked tokens (theme-lock.yaml v1) ─────────────────────────────────
PAPER, PAPER2, INK, YEL = "#ECEAE4", "#F4F2EC", "#1A1C1F", "#F6B800"
STEEL, SOFT, BODY, RULE = "#C9CBC6", "#5E6268", "#33363B", "#CFCDC8"
YEL_TINT = "#F1D98A"
DISP, SANS, MONO = "Bahnschrift SemiBold Condensed", "Calibri", "Consolas"
TOTAL = 14


def px(v: float) -> float:
    """1920-px stage units → inches (13.333 in wide)."""
    return v / 144.0


def resolved_theme() -> dict:
    t = ThemeComposer().compose(layout="standard", seed=23)
    t["colors"].update({"primary": INK, "on-primary": PAPER, "secondary": SOFT, "accent": YEL,
                        "background": PAPER, "foreground": INK, "muted": PAPER2,
                        "muted-foreground": SOFT, "border": RULE})
    t["typography"] = {"heading": DISP, "body": SANS}
    t.setdefault("semantic_roles", {}).update({"background": PAPER, "surface": PAPER2, "ink": INK,
                                               "muted": SOFT, "accent": YEL, "border": RULE,
                                               "data-series-1": INK, "data-series-2": YEL})
    validate_resolved_theme(t)
    blob = json.dumps(t, sort_keys=True, default=str)
    (HERE / "resolved-theme-v1.json").write_text(json.dumps(t, indent=2, default=str), encoding="utf-8")
    print("resolved theme sha256:", hashlib.sha256(blob.encode()).hexdigest())
    return t


# ── Decorative backgrounds (texture + hazard stripes only; no content) ──
def _dots(d, w, h, color, step=56):
    for y in range(28, h, step):
        for x in range(28, w, step):
            d.ellipse([x - 2.6, y - 2.6, x + 2.6, y + 2.6], fill=color)


def _hazard(d, x0, y0, x1, y1, band):
    """Diagonal yellow/black stripes clipped to the box."""
    d.rectangle([x0, y0, x1, y1], fill=INK)
    span = (x1 - x0) + (y1 - y0)
    for k in range(-span, span, band * 2):
        poly = [(x0 + k, y1), (x0 + k + band, y1), (x0 + k + band + (y1 - y0), y0), (x0 + k + (y1 - y0), y0)]
        d.polygon(poly, fill=YEL)


def backgrounds() -> tuple[Path, Path]:
    ASSETS.mkdir(exist_ok=True)
    W, H = 3840, 2160
    content = Image.new("RGB", (W, H), PAPER)
    d = ImageDraw.Draw(content)
    _dots(d, W, H, "#DEDCD6")
    stripe = Image.new("RGB", (44, H))
    _hazard(ImageDraw.Draw(stripe), 0, 0, 43, H - 1, 36)
    content.paste(stripe, (0, 0))
    content.save(ASSETS / "bg-content.png", optimize=True)

    cover = Image.new("RGB", (W, H), INK)
    d = ImageDraw.Draw(cover)
    _dots(d, W, H, "#26292D")
    band = Image.new("RGB", (W, 128))
    _hazard(ImageDraw.Draw(band), 0, 0, W - 1, 127, 68)
    cover.paste(band, (0, H - 128))
    cover.save(ASSETS / "bg-cover.png", optimize=True)
    return ASSETS / "bg-content.png", ASSETS / "bg-cover.png"


def fitted(name: str, w: float, h: float, recolor: dict | None = None) -> Path:
    """Pad a transparent scheme PNG to the target box ratio so cover_image() never crops it."""
    src = Image.open(PNG / f"{name}.png").convert("RGBA")
    if recolor:
        px_ = src.load()
        for y in range(src.height):
            for x in range(src.width):
                r, g, b, a = px_[x, y]
                if a:
                    key = "#%02X%02X%02X" % (r, g, b)
                    if key in recolor:
                        nr, ng, nb = (int(recolor[key][i:i + 2], 16) for i in (1, 3, 5))
                        px_[x, y] = (nr, ng, nb, a)
    ratio = w / h
    sw, sh = src.size
    if sw / sh > ratio:
        cw, ch = sw, round(sw / ratio)
    else:
        cw, ch = round(sh * ratio), sh
    canvas = Image.new("RGBA", (cw, ch), (0, 0, 0, 0))
    canvas.paste(src, ((cw - sw) // 2, (ch - sh) // 2), src)
    out = ASSETS / f"fit-{name}-{int(w)}x{int(h)}{'-inv' if recolor else ''}.png"
    canvas.save(out, optimize=True)
    return out


def scheme(s, name, x, y, w, h, recolor=None):
    return cover_image(s, px(x), px(y), px(w), px(h), str(fitted(name, w, h, recolor)))


# ── Drawing helpers (public pptx_designer helpers) ───────────────────────
def _rgb(c):
    return RGBColor.from_string(c.lstrip("#"))


def t(s, x, y, w, h, txt, size, font=SANS, color=INK, align="left", italic=False, bold=False,
      spacing=None, anchor="top", line=None):
    """Text box at stage px; size in stage px (pt = px / 2). Zero insets for exact placement."""
    box = text(s, px(x), px(y), px(w), px(h), txt, font_size=size / 2, color=color, bold=bold,
               align=align, font_name=font, anchor=anchor)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    run = tf.paragraphs[0].runs[0]
    run.font.italic = italic
    if spacing is not None:
        run.font._rPr.set("spc", str(int(spacing * size / 2 * 100)))
    if line:
        tf.paragraphs[0].line_spacing = line
    return box


def add_run(box, txt, size, font=SANS, color=INK, italic=False, bold=False, new_par=False, line=None,
            space_before=0):
    tf = box.text_frame
    p = tf.add_paragraph() if new_par else tf.paragraphs[-1]
    if new_par:
        p.alignment = tf.paragraphs[0].alignment
        p.space_before = Pt(space_before)
        if line:
            p.line_spacing = line
    r = p.add_run()
    r.text = txt
    r.font.size, r.font.name, r.font.italic, r.font.bold = Pt(size / 2), font, italic, bold
    r.font.color.rgb = _rgb(color)
    return r


def rich(s, x, y, w, h, parts, size, color=BODY, line=1.25):
    """Paragraph made of (text, bold) parts."""
    box = t(s, x, y, w, h, parts[0][0], size, color=color, bold=parts[0][1], line=line)
    for txt, b in parts[1:]:
        add_run(box, txt, size, color=INK if b else color, bold=b)
    return box


def fill(s, x, y, w, h, color, line=None, weight=3):
    shp = rect(s, px(x), px(y), px(w), px(h), color, line=line)
    if line:
        shp.line.width = Pt(weight / 2)
    return shp


def hline(s, x, y, w, weight=2, color=INK):
    return rect(s, px(x), px(y), px(w), px(weight), color)


def new_slide(prs, bg, n, tag, section, dark=False):
    s = add_slide(prs)
    cover_image(s, 0, 0, 13.333, 7.5, str(bg))
    left = 112 if dark else 150
    tw = 22 + len(tag) * 12
    fill(s, left, 34, tw, 32, INK if not dark else YEL)
    t(s, left, 34, tw, 32, tag, 19, MONO, YEL if not dark else INK, align="center", bold=True, anchor="middle")
    t(s, 560, 38, 800, 26, section.upper(), 19, MONO, "#9A9EA4" if dark else SOFT, align="center", spacing=0.1)
    t(s, 1508, 38, 300, 26, f"{n:02d} / {TOTAL}", 19, MONO, PAPER if dark else INK, align="right", bold=True)
    return s


def headline(s, txt, y=104, w=1660):
    # Conservative auto-size: keeps long titles on one line even with a wider fallback font.
    size = min(104, int(w / (0.6 * len(txt))))
    t(s, 150, y + (104 - size), w, 120, txt.upper(), size, DISP, INK, bold=False, line=0.9)
    fill(s, 150, y + 128, 240, 14, YEL)


def foot(s, txt):
    t(s, 150, 1022, 1660, 26, txt, 17, MONO, SOFT)


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


def ledger(s, x, y, widths, header, rows, row_h, size=25, focus=None, first_display=True, mono_cols=()):
    """Table drawn from native text boxes and hairlines (no default table styling)."""
    W = sum(widths)
    fill(s, x, y, W, 54, INK)
    cx = x
    for wcol, htxt in zip(widths, header):
        t(s, cx + 18, y, wcol - 24, 54, htxt.upper(), 18, MONO, PAPER, bold=True, anchor="middle", spacing=0.08)
        cx += wcol
    ry = y + 54
    for i, row in enumerate(rows):
        if focus == i:
            fill(s, x, ry, W, row_h, YEL_TINT)
            fill(s, x, ry, 8, row_h, YEL)
        cx = x
        for j, (wcol, cell) in enumerate(zip(widths, row)):
            if j == 0 and first_display:
                t(s, cx + 18, ry, wcol - 24, row_h, cell.upper(), 28, DISP, INK, anchor="middle")
            else:
                font = MONO if j in mono_cols else SANS
                bold = isinstance(cell, tuple)
                val = cell[0] if bold else cell
                t(s, cx + 18, ry, wcol - 28, row_h, val, size if font == SANS else size - 2, font, INK if (bold or j == 0) else BODY,
                  bold=bold, anchor="middle", line=1.1)
            cx += wcol
        ry += row_h
        hline(s, x, ry - 1.5, W, 1.5, RULE)
    return ry


# ── Slides ───────────────────────────────────────────────────────────────
def s01_cover(prs, bg):
    s = new_slide(prs, bg, 1, "PBL · TA1", "Įvadas į specialybę · Mechanikos inžinerija", dark=True)
    t(s, 112, 196, 900, 30, "TARPINIS ATSISKAITYMAS NR. 1", 20, MONO, YEL, bold=True, spacing=0.14)
    t(s, 104, 250, 1150, 230, "DOMKRATAS", 190, DISP, YEL, line=0.85)
    t(s, 112, 520, 900, 120, "Nuo sraigto iki elektros: raida, problemos ir patobulinimo kryptis", 40, SANS, PAPER, line=1.2)
    scheme(s, "scissor", 1260, 250, 560, 480, recolor={INK: "#ECEAE4", "#F4F2EC": INK, SOFT: "#8B8F95"})
    for i, (k, v) in enumerate([("GRUPĖ", "[Vardas Pavardė ×5–6]"), ("DĖSTYTOJAS", "G. Viselga"), ("DATA", "2026 m. spalis")]):
        x = 112 + i * 330
        t(s, x, 880, 320, 24, k, 17, MONO, "#8B8F95", bold=True, spacing=0.14)
        t(s, x, 910, 320, 34, v, 24, SANS, "#C9CBC6")
    notes(s, "Laba diena. Mūsų PBL tema – domkratas. Per 4 minutes parodysime, kaip domkratai vystėsi, "
             "kokias problemas turi šiandien ir kaip siūlome patobulinti automobilio žirklinį domkratą.")


def s02_why(prs, bg):
    s = new_slide(prs, bg, 2, "PBL 1", "Srities pasirinkimas")
    headline(s, "Kodėl domkratas?")
    items = [("01", "Jį turi beveik kiekvienas automobilis",
              "Paprasčiausias kėlimo mechanizmas, su kuriuo susiduria kiekvienas vairuotojas – ypač pradurtos padangos atveju."),
             ("02", "Visas mūsų dalykas – viename įrenginyje",
              "Pavaros (sraigtas, krumpliastiebis, hidraulika), medžiagos, sauga, ergonomika ir kaina."),
             ("03", "Klaidos kaina – žmogaus sauga",
              "Domkratas laiko šimtus kilogramų virš žmogaus rankų, todėl kiekvienas sprendimas turi būti pagrįstas.")]
    y = 300
    for n, h, p in items:
        hline(s, 150, y, 1000, 2.5)
        t(s, 150, y + 20, 80, 60, n, 54, DISP, INK)
        t(s, 245, y + 20, 900, 44, h, 34, SANS, INK, bold=True)
        t(s, 245, y + 70, 900, 90, p, 27, SANS, BODY, line=1.2)
        y += 210
    fill(s, 1220, 300, 600, 470, INK)
    stripe = ASSETS / "stripe-thin.png"
    if not stripe.exists():
        im = Image.new("RGB", (1200, 28))
        _hazard(ImageDraw.Draw(im), 0, 0, 1199, 27, 24)
        im.save(stripe)
    cover_image(s, px(1220), px(300), px(600), px(14), str(stripe))
    t(s, 1272, 350, 520, 26, "VIENAM AUTOMOBILIO KAMPUI TENKA", 19, MONO, "#8B8F95", bold=True, spacing=0.1)
    t(s, 1264, 384, 540, 170, "≈0,5 t", 160, DISP, YEL)
    t(s, 1272, 566, 500, 80, "kai automobilio masė ~1,5 t: 0,4–0,5 t (masė ÷ 4 + atsarga dėl svorio centro).", 26, SANS, "#D9DBD6", line=1.2)
    t(s, 1272, 676, 520, 70, "autoservisai · ratų keitimas · statyba · geležinkelis · gelbėtojai", 20, MONO, PAPER, line=1.3)
    notes(s, "Pasirinkome domkratą, nes jį turi beveik kiekvienas automobilis. Jame susitinka viskas, ką mokomės: "
             "pavaros, medžiagos, sauga ir ergonomika. Vienam automobilio kampui tenka apie pusę tonos.")


def s03_principles(prs, bg):
    s = new_slide(prs, bg, 3, "PBL 1", "Veikimo principas")
    headline(s, "Kaip ranka pakelia pusę tonos?")
    for i, (img, ttl, f, parts) in enumerate([
        ("incline", "Mechaninis: sraigtas", "F = Q·p / (2πR·η)",
         [("Pvz.: Q = 5 kN, žingsnis p = 4 mm, rankena R = 0,15 m, η ≈ 0,35 → ", False), ("F ≈ 60 N", True),
          (". Laimime jėgą, bet pralaimime kelią: daug apsisukimų.", False)]),
        ("pascal", "Hidraulinis: Paskalio dėsnis", "F₂ = F₁·(D₂ / D₁)²",
         [("Pvz.: D₁ = 10 mm, D₂ = 40 mm → jėga ×16; su svirtimi 1:10 → ", False), ("×160", True),
          (". 200 N rankos jėga → ~32 kN.", False)])]):
        x = 150 + i * 860
        hline(s, x, 300, 800, 6)
        scheme(s, img, x, 326, 800, 330)
        t(s, x, 676, 800, 60, ttl.upper(), 50, DISP, INK)
        fill(s, x, 750, 470, 58, YEL)
        t(s, x + 20, 750, 450, 58, f, 32, MONO, INK, bold=True, anchor="middle")
        rich(s, x, 830, 790, 110, parts, 26)
    foot(s, "Šaltinis: Budynas & Nisbett (2020), sraigtinių pavarų skyrius. Skaičiai – mūsų pavyzdiniai įverčiai.")
    notes(s, "Visi domkratai remiasi dviem principais. Sraigtas – tai aplink cilindrą apvyniota nuožulni plokštuma: "
             "maža jėga, bet daug apsisukimų. Hidraulikoje Paskalio dėsnis: slėgis vienodas, todėl didelis stūmoklis duoda didelę jėgą.")


def s04_timeline(prs, bg):
    s = new_slide(prs, bg, 4, "PBL 2", "Istorinė raida")
    headline(s, "Nuo sverto iki elektros")
    ev = [("III a. pr. Kr.", "Svertas ir pleištas", "Archimedas aprašo sverto principą – jėgos laimėjimo pagrindą."),
          ("XV–XVI a.", "Sraigtinis kėliklis", "Sraigtas su veržle; Leonardo da Vinci eskizuose – kėliklis su krumpliaratine pavara."),
          ("1795", "Hidraulinis presas", "J. Bramah patentuoja presą, taikantį Paskalio dėsnį."),
          ("1851", "Nešiojamas hidraulinis", "R. Dudgeon (JAV) – kompaktiškas „butelinis“ domkratas."),
          ("XX a.", "Automobilių era", "1905 m. „Hi-Lift“ krumpliastiebinis; vėliau žirklinis tampa automobilio komplektu, servisuose – vežimėliniai."),
          ("XXI a.", "Elektra ir pneumatika", "12 V elektriniai, elektrohidrauliniai, pneumatinės pagalvės.")]
    widths = [300, 260, 230, 260, 300, 310]
    hline(s, 150, 482, 1660, 6)
    x = 150
    for (yr, h, p), w in zip(ev, widths):
        t(s, x, 400, w, 70, yr, 54, DISP, INK)
        d = rect(s, px(x + 2), px(470), px(28), px(28), YEL if yr != "XXI a." else INK, line=INK)
        d.rotation = 45
        d.line.width = Pt(2.5)
        t(s, x, 528, w - 24, 70, h, 27, SANS, INK, bold=True, line=1.1)
        t(s, x, 606, w - 24, 230, p, 23, SANS, BODY, line=1.2)
        x += w
    labels = ["JĖGOS LAIMĖJIMAS", "DIDESNĖ GALIA", "KOMPAKTIŠKUMAS", "SAUGA IR ERGONOMIKA"]
    for i, lab in enumerate(labels):
        bx = 150 + i * 415
        fill(s, bx, 900, 417, 60, YEL if i == 3 else PAPER2, line=INK, weight=2.5)
        t(s, bx + 22, 900, 390, 60, lab, 22, MONO, INK, bold=True, anchor="middle", spacing=0.08)
    foot(s, "Šaltiniai: Encyclopaedia Britannica (n.d.); Hi-Lift Jack Company (n.d.). Datas patikslinti rašto darbe.")
    notes(s, "Raida: nuo Archimedo sverto, per sraigtinius kėliklius, Bramah hidraulinį presą ir Dudgeon nešiojamą domkratą "
             "iki automobilių eros ir šiandieninių elektrinių. Tendencija – nuo jėgos laimėjimo link saugos ir ergonomikos.")


def s05_types(prs, bg):
    s = new_slide(prs, bg, 5, "PBL 2", "Principinės schemos")
    headline(s, "Penki domkratų tipai")
    cards = [("screw", "A · MECHANINIS", "Sraigtinis", "Sukama rankena suka sraigtą, kuris išsuka stūmoklį į viršų.", "Savistabdis, paprastas"),
             ("rack", "B · MECHANINIS", "Krumplia-\nstiebinis", "Svirtis su skląsčiu kopia krumpliastiebiu.", "Didelis aukštis"),
             ("scissor", "C · MECHANINIS", "Žirklinis", "Horizontalus sraigtas suartina šonines ašis – rombas kyla.", "Kompaktiškas, pigus"),
             ("bottle", "D · HIDRAULINIS", "Butelinis", "Maža pompa spaudžia alyvą po dideliu stūmokliu.", "Didelė jėga"),
             ("trolley", "E · HIDRAULINIS", "Vežimėlinis", "Hidrocilindras kelia ilgą svirtį; važiuoja ant ratukų.", "Greitas, stabilus")]
    w, gap = 314, 22
    for i, (img, idt, h, p, pr) in enumerate(cards):
        x = 150 + i * (w + gap)
        fill(s, x, 296, w, 690, YEL if img == "scissor" else PAPER2, line=INK, weight=3)
        scheme(s, img, x + 14, 312, w - 28, 340)
        t(s, x + 20, 664, w - 40, 22, idt, 17, MONO, INK if img == "scissor" else SOFT, bold=True, spacing=0.1)
        lines = h.upper().split("\n")
        box = t(s, x + 20, 694, w - 40, 110, lines[0], 44, DISP, INK, line=0.95)
        for ln in lines[1:]:
            add_run(box, ln, 44, DISP, INK, new_par=True, line=0.95)
        if len(lines) > 1:
            continue_y = 46
        else:
            continue_y = 0
        t(s, x + 20, 790 + continue_y, w - 40, 130, p, 24, SANS, BODY, line=1.2)
        t(s, x + 20, 930, w - 40, 30, pr, 20, MONO, INK, bold=True)
    notes(s, "Penki pagrindiniai tipai: trys mechaniniai – sraigtinis, krumpliastiebinis, žirklinis – ir du hidrauliniai – "
             "butelinis ir vežimėlinis. Toliau gilinsimės į žirklinį, nes jis yra kiekvieno automobilio komplekte.")


def s06_compare(prs, bg):
    s = new_slide(prs, bg, 6, "PBL 2", "Palyginimas")
    headline(s, "Kuris ką daro geriausiai")
    rows = [["Sraigtinis", "2–20 t", "~100–300 mm", "3–15 kg", "Savistabdis, patikimas", "Lėtas, mažas aukštis"],
            ["Krumpliastiebinis", "~1,5–2 t", "iki ~1 m", "~10–15 kg", "Labai didelis aukštis", "Nestabilus, rankenos atatranka"],
            ["Žirklinis", "0,8–2 t", "~100–400 mm", "2–5 kg", "Kompaktiškas, pigus, savistabdis", "Lėtas, sunku sukti, siauras pagrindas"],
            ["Butelinis", "2–50 t", "~150–450 mm", "3–15 kg", "Didžiausia jėga mažame tūryje", "Aukštas pradinis aukštis, nutekėjimai"],
            ["Vežimėlinis", "2–3,5 t", "~80–500 mm", "15–40 kg", "Greitas, stabilus", "Sunkus, nenešiojamas, brangus"]]
    ledger(s, 150, 296, [330, 200, 220, 180, 360, 370], ["Tipas", "Keliamoji galia", "Kėlimo aukštis", "Masė", "Stiprybė", "Silpnybė"],
           rows, 112, size=25, focus=2, mono_cols=(1, 2, 3))
    foot(s, "Tipinės rinkoje parduodamų gaminių vertės (apytikslės, iš gamintojų katalogų). Tikslius skaičius ir šaltinius pateikti rašto darbe.")
    notes(s, "Palyginimas rodo kompromisą: hidrauliniai stipresni, bet sunkesni ar aukštesni; žirklinis – lengviausias ir kompaktiškiausias, "
             "bet lėtas, sunkiai sukamas ir su siauru pagrindu.")


def s07_verdict(prs, bg):
    s = new_slide(prs, bg, 7, "PBL 2", "Kritinis vertinimas")
    headline(s, "Kas pasiteisino, o kas ne")
    cols = [("✓", "Pasiteisino", PAPER2, INK, BODY,
             [("Savistabdis trapecinis sriegis", " – krovinys nenukrenta net paleidus rankeną."),
              ("Hidraulika", " – didžiausia jėga mažiausiame tūryje."),
              ("Žirklinė geometrija", " – telpa po žemu automobiliu ir bagažinėje."),
              ("Platus pagrindas su ratukais", " (vežimėlinis) – stabilumas servise.")]),
            ("✗", "Nepasiteisino", INK, PAPER, "#D9DBD6",
             [("Rankenos atatranka", " krumpliastiebiniuose – dažna traumų priežastis."),
              ("Siauras plonos skardos pagrindas", " – slysta, smenga į gruntą ar karštą asfaltą."),
              ("Hidraulika be perkrovos vožtuvo", " – trūkus sandarikliui krovinys gali nusileisti."),
              ("Ilgas rankinis sukimas", " – žmogus ilgai būna šalia važiuojamosios dalies.")])]
    for i, (ic, h, bgc, hc, bc, items) in enumerate(cols):
        x = 150 + i * 830
        fill(s, x, 296, 830, 660, bgc, line=INK, weight=4)
        fill(s, x + 48, 340, 60, 60, YEL)
        t(s, x + 48, 340, 60, 60, ic, 38, SANS, INK, bold=True, align="center", anchor="middle")
        t(s, x + 128, 338, 600, 66, h.upper(), 54, DISP, hc, anchor="middle")
        y = 440
        for b, rest in items:
            fill(s, x + 48, y + 14, 12, 12, INK if i == 0 else YEL)
            box = t(s, x + 76, y, 720, 110, b, 28, SANS, hc, bold=True, line=1.2)
            add_run(box, rest, 28, SANS, bc)
            y += 122
    notes(s, "Pasiteisino savistabdis sriegis, hidraulika ir žirklinė geometrija. Nepasiteisino rankenos atatranka, siauri pagrindai, "
             "hidraulika be apsaugos ir ilgas rankinis sukimas – būtent šias silpnybes ir taikysime.")


def s08_problem(prs, bg):
    s = new_slide(prs, bg, 8, "PBL 1", "Šiuolaikinė problema")
    headline(s, "Žirklinis domkratas kelkraštyje")
    fill(s, 150, 296, 640, 670, PAPER2, line=INK, weight=3)
    scheme(s, "scissor", 176, 316, 588, 420)
    fill(s, 176, 752, 380, 56, YEL)
    t(s, 192, 752, 360, 56, "F sraigto = Q / tg θ", 30, MONO, INK, bold=True, anchor="middle")
    rich(s, 176, 824, 590, 130, [("Kuo žemiau domkratas, tuo mažesnis kampas θ ir tuo didesnė jėga sraigte: kai θ ≈ 15°, ji beveik ", False),
                                  ("4 kartus", True), (" viršija krovinį.", False)], 23)
    cards = [("60–100 aps.", "Lėta", "Tiek rankenos apsisukimų reikia pakelti iki darbinio aukščio (~2–4 min)."),
             ("×4", "Sunku pradžioje", "Apkrovos piką rankenoje jaučia senyvi ir silpnesni vairuotojai."),
             ("~100 cm²", "Mažas pagrindas", "Ant minkšto grunto ar nuolydžio domkratas gali pasvirti."),
             ("0,5 m", "Ergonomika", "Žmogus dirba pasilenkęs, rankena kliūna už kėbulo ar rato.")]
    for i, (k, h, p) in enumerate(cards):
        x = 860 + (i % 2) * 490
        y = 296 + (i // 2) * 340
        hline(s, x, y, 460, 6)
        t(s, x, y + 18, 460, 90, k, 76, DISP, INK)
        t(s, x, y + 120, 460, 40, h, 28, SANS, INK, bold=True)
        t(s, x, y + 166, 450, 110, p, 24, SANS, BODY, line=1.2)
    foot(s, "Apsisukimų skaičius – mūsų įvertinimas: svirtis 200 mm, kėlimas 100→350 mm, sriegio žingsnis 2–3 mm.")
    notes(s, "Šiandienos problema – žirklinis domkratas automobilio komplekte. Jis lėtas: 60–100 apsisukimų. Pradžioje sunkiausia, "
             "nes jėga sraigte lygi Q padalinta iš tangento θ. Pagrindas mažas, o žmogus dirba pasilenkęs prie kelio.")


def s09_task(prs, bg):
    s = new_slide(prs, bg, 9, "PBL 1", "Uždavinio formulavimas")
    headline(s, "Ką sprendžiame")
    fill(s, 150, 296, 930, 420, YEL)
    t(s, 206, 336, 820, 26, "INŽINERINIS UŽDAVINYS", 19, MONO, INK, bold=True, spacing=0.14)
    box = t(s, 206, 380, 830, 320, "Patobulinti automobilio žirklinį domkratą taip, kad 1,5 t automobilį jis pakeltų ", 40, SANS, INK, line=1.2)
    add_run(box, "greičiau nei per 60 s", 40, SANS, INK, bold=True)
    add_run(box, ", be didelių fizinių pastangų ir stabiliai ant nelygaus pagrindo – išlaikant kompaktiškumą ir savistabdį saugumą.", 40, SANS, INK)
    t(s, 150, 750, 930, 160, "Kodėl ne hidraulinis? Butelinis domkratas per aukštas po žemu automobiliu, o vežimėlinis netelpa bagažinėje. "
                             "Žirklinė geometrija lieka geriausiu pagrindu – tobuliname jos silpnąsias vietas.", 27, SANS, BODY, line=1.3)
    rows = [["Keliamoji galia", "≥ 1500 kg"], ["Kėlimo diapazonas", "100–400 mm"], ["Kėlimo laikas", "≤ 60 s"],
            ["Rankos jėga", "0 N (el.) / ≤ 150 N"], ["Pagrindo plotas", "≈ 2× didesnis"], ["Masė", "≤ 4,5 kg"],
            ["Maitinimas", "12 V lizdas, ≤ 10 A"]]
    ledger(s, 1150, 296, [340, 320], ["Reikalavimas", "Tikslas"], rows, 86, size=26, first_display=False, mono_cols=(1,))
    notes(s, "Mūsų inžinerinis uždavinys: pakelti 1,5 t automobilį greičiau nei per minutę, be fizinių pastangų ir stabiliai. "
             "Lentelėje – konkretūs tikslai, pagal kuriuos vertinsime sprendimą.")


def s10_concept(prs, bg):
    s = new_slide(prs, bg, 10, "PBL 2", "Sprendimo idėja")
    headline(s, "Elektrinis žirklinis domkratas")
    fill(s, 150, 280, 1100, 640, PAPER2, line=INK, weight=3)
    scheme(s, "modern", 170, 300, 1060, 600)
    legend = [("12 V DC variklis", " (~100 W) iš automobilio lizdo"),
              ("Planetinis reduktorius", " – didelis sukimo momentas, kompaktiškas"),
              ("Trapecinis sraigtas Tr18×4", " – greitesnis, savistabdis"),
              ("Bronzinė veržlė", " – maža trintis, mažiau dyla"),
              ("Svirtys iš S420MC", " – stipresnės be papildomos masės"),
              ("Platesnis pagrindas + guminis padas", " – neslysta, nesmenga"),
              ("Avarinis šešiakampis", " – rankinis sukimas be elektros"),
              ("Galiniai jungikliai, srovės ribotuvas", " ir guminė atrama")]
    y = 284
    for i, (b, rest) in enumerate(legend):
        oval(s, px(1290), px(y), px(42), px(42), YEL, line=INK).line.width = Pt(1.5)
        t(s, 1290, y, 42, 42, str(i + 1), 20, MONO, INK, bold=True, align="center", anchor="middle")
        box = t(s, 1350, y + 4, 470, 74, b, 24, SANS, INK, bold=True, line=1.15)
        add_run(box, rest, 24, SANS, BODY)
        y += 82
    notes(s, "Mūsų idėja – elektrinis žirklinis domkratas: 12 V variklis su planetiniu reduktoriumi suka trapecinį sraigtą su bronzine veržle. "
             "Stipresnės svirtys, platesnis guminis pagrindas, avarinis rankinis sukimas ir apsaugos nuo perkrovos.")


def s11_materials(prs, bg):
    s = new_slide(prs, bg, 11, "PBL 2", "Keičiami elementai ir medžiagos")
    headline(s, "Ką keičiame ir kodėl")
    rows = [["Pavara", "Rankinė rankena", ("12 V DC variklis + planetinis reduktorius",), "Greitis, nereikia fizinės jėgos"],
            ["Sraigtas", "Metrinis sriegis, paprastas plienas", ("Tr18×4, plienas 42CrMo4 (grūdintas)",), "Savistabdis (ψ < φ′), atsparus dilimui ir lenkimui"],
            ["Veržlė", "Plieninė", ("Bronza CuSn12",), "Mažesnė trintis su plienu, nesukimba"],
            ["Svirtys", "S235 štampuota skarda", ("S420MC to paties storio",), "Takumo riba 420 vs 235 MPa (~1,8×)"],
            ["Pagrindas", "~130×80 mm, plona skarda", ("~200×120 mm + guminis padas",), "~2× didesnis plotas, neslysta"],
            ["Korpusas", "—", ("PA6-GF30 (stiklo pluoštu armuotas)",), "Lengvas, smūgiams ir alyvai atsparus"]]
    ledger(s, 150, 296, [250, 420, 500, 490], ["Elementas", "Dabar", "Siūloma", "Pagrindimas"], rows, 104, size=25)
    foot(s, "Medžiagų savybės: EN 10025-2 (S235), EN 10149-2 (S420MC), EN 10083-3 (42CrMo4). Galutinis parinkimas ir skaičiavimai – TA2.")
    notes(s, "Konkrečiai keičiame šešis elementus. Svarbiausia: sraigtas iš grūdinto 42CrMo4 plieno su bronzine veržle – savistabdis ir mažai dyla; "
             "svirtys iš aukšto stiprumo S420MC plieno; pagrindas dvigubai didesnis.")


def s12_params(prs, bg):
    s = new_slide(prs, bg, 12, "PBL 2", "Parametrai")
    headline(s, "Prieš ir po")
    rows = [("Kėlimo laikas", "~2–4 min", "≤ 60 s", True), ("Rankos jėga", "~100–200 N", "0 N", True),
            ("Apsisukimai", "60–100", "0", True), ("Pagrindo plotas", "~100 cm²", "~240 cm²", True),
            ("Masė", "~3 kg", "~4,5 kg", False), ("Kaina", "~20–30 €", "~60–80 €", False)]
    widths = [300, 260, 260, 240]
    x0, y0 = 150, 296
    fill(s, x0, y0, sum(widths), 54, INK)
    for j, htxt in enumerate(["Parametras", "Įprastas", "Patobulintas", ""]):
        t(s, x0 + sum(widths[:j]) + 18, y0, widths[j] - 24, 54, htxt.upper(), 18, MONO, PAPER, bold=True, anchor="middle", spacing=0.08)
    y = y0 + 54
    for name, old, new, good in rows:
        t(s, x0 + 18, y, 280, 100, name, 26, SANS, INK, anchor="middle")
        t(s, x0 + 318, y, 240, 100, old, 25, MONO, BODY, anchor="middle")
        t(s, x0 + 578, y, 240, 100, new, 25, MONO, INK, bold=True, anchor="middle")
        tag_w = 130 if good else 190
        fill(s, x0 + 838, y + 34, tag_w, 32, YEL if good else PAPER, line=None if good else INK, weight=2.5)
        t(s, x0 + 838, y + 34, tag_w, 32, "GERIAU" if good else "KOMPROMISAS", 16, MONO, INK, bold=True, align="center", anchor="middle")
        y += 100
        hline(s, x0, y - 1.5, sum(widths), 1.5, RULE)
    fill(s, 1260, 296, 560, 660, INK)
    t(s, 1302, 336, 480, 26, "PRELIMINARUS ĮVERTINIMAS", 19, MONO, YEL, bold=True, spacing=0.12)
    calc = [("P ≈ Q·h / (t·η)", "5000 N · 0,25 m / (45 s · 0,3) ≈ 93 W → ~8 A iš 12 V lizdo (saugiklis 15 A)."),
            ("ψ = 4,5° < φ′ ≈ 5,9°", "Tr18×4 su bronzine veržle (μ ≈ 0,1) – sriegis savistabdis: krovinys nenusileidžia dingus elektrai."),
            ("η sraigto ≈ 0,4", "Visos pavaros η ≈ 0,4 · 0,9 · 0,85 ≈ 0,3.")]
    y = 390
    for i, (eq, p) in enumerate(calc):
        if i:
            hline(s, 1302, y - 22, 476, 1.5, "#3A3D42")
        t(s, 1302, y, 476, 44, eq, 30, MONO, YEL, bold=True)
        t(s, 1302, y + 50, 476, 120, p, 23, SANS, "#C9CBC6", line=1.25)
        y += 190
    notes(s, "Palyginus: kėlimas trumpėja nuo kelių minučių iki minutės, rankos jėga lygi nuliui, pagrindas dvigubai didesnis. "
             "Kompromisai – masė ir kaina. Preliminariai variklio galia apie 100 W, o sriegis išlieka savistabdis.")


def s13_conclusions(prs, bg):
    s = new_slide(prs, bg, 13, "TA1", "Išvados")
    headline(s, "Išvados ir kas toliau")
    items = ["Domkratų raida: nuo jėgos laimėjimo (svertas, sraigtas) per hidrauliką iki ergonomikos ir automatizacijos.",
             "Žirklinis domkratas pasiteisino kompaktiškumu ir savistabdžiu sriegiu, bet atsilieka greičiu, ergonomika ir stabilumu.",
             "Siūlome: 12 V elektrinę pavarą, Tr18×4 sraigtą su bronzine veržle, stipresnes svirtis ir platesnį pagrindą.",
             "Kompromisas – šiek tiek didesnė masė ir kaina, kurias pateisina greitis ir sauga."]
    y = 296
    for i, it in enumerate(items):
        hline(s, 150, y, 1000, 2.5)
        t(s, 150, y + 22, 80, 50, f"0{i + 1}", 42, DISP, INK)
        t(s, 240, y + 22, 910, 120, it, 30, SANS, INK, line=1.25)
        y += 165
    fill(s, 1220, 296, 600, 400, PAPER2, line=INK, weight=4)
    t(s, 1258, 330, 520, 56, "TA2 PLANAS", 46, DISP, INK)
    nxt = ["Naujausių patentų ir straipsnių analizė", "Sraigto stiprumo, klupumo ir veržlės dilimo skaičiavimai",
           "Variklio ir reduktoriaus parinkimas", "Atitiktis EN 1494, sauga ir aplinkosauga"]
    yy = 404
    for n in nxt:
        t(s, 1258, yy, 30, 40, "→", 26, MONO, INK)
        t(s, 1290, yy, 500, 70, n, 26, SANS, INK, line=1.15)
        yy += 70
    fill(s, 1220, 730, 600, 170, YEL)
    t(s, 1258, 730, 540, 170, "KLAUSIMAI?", 96, DISP, INK, anchor="middle")
    notes(s, "Apibendrinant: domkratai vystėsi nuo jėgos laimėjimo iki saugos ir ergonomikos. Žirklinį domkratą siūlome elektrifikuoti ir sustiprinti. "
             "TA2 etape atliksime skaičiavimus, patentų analizę ir atitikties standartams vertinimą. Ačiū, laukiame klausimų.")


def s14_refs(prs, bg):
    s = new_slide(prs, bg, 14, "APA 7", "Šaltiniai")
    headline(s, "Šaltiniai")
    refs = [[("Budynas, R. G., & Nisbett, J. K. (2020). ", False), ("Shigley's mechanical engineering design", True), (" (11th ed.). McGraw-Hill Education.", False)],
            [("Childs, P. R. N. (2014). ", False), ("Mechanical design engineering handbook", True), (". Butterworth-Heinemann.", False)],
            [("European Committee for Standardization. (2008). ", False), ("Mobile or movable jacks and associated lifting equipment", True), (" (EN 1494:2000+A1:2008).", False)],
            [("European Committee for Standardization. (2019). ", False), ("Hot rolled products of structural steels – Part 2", True), (" (EN 10025-2:2019).", False)],
            [("Europos Parlamentas ir Taryba. (2023). ", False), ("Reglamentas (ES) 2023/1230 dėl mašinų", True), (". Europos Sąjungos oficialusis leidinys.", False)],
            [("Encyclopaedia Britannica. (n.d.). ", False), ("Jack", True), (". Britannica. Retrieved [data], from https://www.britannica.com", False)],
            [("Hi-Lift Jack Company. (n.d.). ", False), ("Our history", True), (". Retrieved [data], from https://www.hi-lift.com", False)]]
    for i, parts in enumerate(refs):
        x = 150 + (i // 4) * 860
        y = 300 + (i % 4) * 150
        box = t(s, x, y, 800, 140, parts[0][0], 27, SANS, INK, line=1.3)
        for txt, it in parts[1:]:
            add_run(box, txt, 27, SANS, INK, italic=it)
    hline(s, 150, 930, 1660, 1.5, RULE)
    t(s, 150, 946, 1660, 60, "Pastaba: prieš atsiskaitymą patikrinkite kiekvieną šaltinį ir įrašykite peržiūros datas; čia pateikiami tik pristatyme minimi šaltiniai.",
      21, MONO, SOFT, line=1.3)
    notes(s, "Šaltiniai pateikti APA 7 stiliumi.")


def main():
    theme = resolved_theme()
    bg_content, bg_cover = backgrounds()
    prs = Presentation(theme=theme, strict_theme=True)
    s01_cover(prs, bg_cover)
    for fn in (s02_why, s03_principles, s04_timeline, s05_types, s06_compare, s07_verdict, s08_problem,
               s09_task, s10_concept, s11_materials, s12_params, s13_conclusions, s14_refs):
        fn(prs, bg_content)
    prs.save(str(OUT))
    print("saved", OUT)


if __name__ == "__main__":
    main()
