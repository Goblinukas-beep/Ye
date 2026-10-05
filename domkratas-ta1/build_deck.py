"""Paprastas PBL TA1 pristatymas „Domkratas“ (10 skaidrių, be piešinių).

Paleidimas:  python3 build_deck.py   → domkratas-ta1.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "domkratas-ta1.pptx"

INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x59, 0x59, 0x59)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
LIGHT = RGBColor(0xDE, 0xEA, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 10


def textbox(slide, x, y, w, h, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.paragraphs[0].alignment = align
    return tf, (size, color, bold)


def para(tf, style, txt, first=False, bullet=False, space=6, size=None, bold=None, color=None):
    s, c, b = style
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space)
    run = p.add_run()
    run.text = ("• " if bullet else "") + txt
    run.font.name, run.font.size = FONT, Pt(size or s)
    run.font.bold = b if bold is None else bold
    run.font.color.rgb = color or c
    return p


def add(p, txt, size, bold=False, italic=False, color=INK):
    r = p.add_run()
    r.text = txt
    r.font.name, r.font.size, r.font.bold, r.font.italic = FONT, Pt(size), bold, italic
    r.font.color.rgb = color
    return r


def bullets(slide, x, y, w, h, items, size=20, space=10):
    """items: (paryškinta dalis, tekstas) arba (paryškinta dalis, tekstas, šaltinis)."""
    tf, st = textbox(slide, x, y, w, h, size)
    for i, it in enumerate(items):
        p = para(tf, st, it[0], first=i == 0, bullet=True, space=space, bold=True)
        add(p, it[1], size)
        if len(it) > 2:
            add(p, f" ({it[2]})", size - 4, color=GREY)
    return tf


def rect(slide, x, y, w, h, color):
    shp = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def new_slide(n, title, notes=""):
    s = prs.slides.add_slide(BLANK)
    tf, st = textbox(s, 0.6, 0.35, 12.1, 0.9, 32, ACCENT, True, anchor=MSO_ANCHOR.BOTTOM)
    para(tf, st, title, first=True)
    rect(s, 0.65, 1.32, 12.0, 0.03, ACCENT)
    tf, st = textbox(s, 11.7, 6.95, 1.1, 0.4, 12, GREY, align=PP_ALIGN.RIGHT)
    para(tf, st, f"{n} / {TOTAL}", first=True)
    tf, st = textbox(s, 0.6, 6.95, 8, 0.4, 12, GREY)
    para(tf, st, "PBL · Domkratas", first=True)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def source_line(slide, txt, y=6.45):
    tf, st = textbox(slide, 0.7, y, 11.9, 0.5, 13, GREY)
    para(tf, st, "Šaltiniai: " + txt, first=True)


def table(slide, x, y, w, col_w, rows, size=15, row_h=0.5, highlight=None):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shape.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if i == 0 else (LIGHT if i == highlight else WHITE)
            p = cell.text_frame.paragraphs[0]
            add(p, val, size, bold=(i == 0 or j == 0), color=WHITE if i == 0 else INK)
    return tbl


# ── 1. Titulinis ────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
rect(s, 0, 2.2, 13.333, 2.4, ACCENT)
tf, st = textbox(s, 0.8, 2.35, 11.7, 1.2, 48, WHITE, True, anchor=MSO_ANCHOR.BOTTOM)
para(tf, st, "Domkratas", first=True)
tf, st = textbox(s, 0.8, 3.55, 11.7, 0.9, 24, WHITE)
para(tf, st, "Raida, problemos ir patobulinimo kryptis", first=True)
tf, st = textbox(s, 0.8, 5.0, 11.7, 1.8, 18, GREY)
para(tf, st, "PBL užduotis · Tarpinis atsiskaitymas Nr. 1", first=True, space=4)
para(tf, st, "Dalykas: Įvadas į specialybę · Mechanikos inžinerija", space=4)
para(tf, st, "Grupė: [Vardas Pavardė, …] · Dėstytojas: G. Viselga · 2026 m.", space=4)
s.notes_slide.notes_text_frame.text = (
    "Laba diena. Mūsų tema – domkratas. Parodysime, kaip domkratai vystėsi, kokių problemų jie turi šiandien "
    "ir kaip siūlome patobulinti automobilio žirklinį domkratą.")

# ── 2. Srities pasirinkimas ─────────────────────────────────────────────
s = new_slide(2, "Kodėl pasirinkome domkratą?",
              "Žirklinį domkratą gamintojai deda į naujus automobilius, todėl su juo susiduria beveik kiekvienas vairuotojas. "
              "Domkratas veikia pagal sverto arba Paskalio dėsnį: nedidele jėga pakeliamas didelis krovinys.")
bullets(s, 0.7, 1.7, 12, 4.6, [
    ("Kasdienis įrenginys", " – žirklinį domkratą gamintojai dažnai deda į naujus automobilius.", "CJ Pony Parts, n.d."),
    ("Platus taikymas", " – automobilių remontas, ūkiai, statyba, geležinkelis.", "Hi-Lift Jack Company, n.d.-a; Wikipedia, n.d.-d"),
    ("Daug mechanikos vienoje vietoje", " – sraigtinė, krumpliastiebinė ir hidraulinė pavaros, medžiagos, sauga."),
    ("Mechaninis principas", " – svertas ir sraigtas: maža jėga ilgu keliu pakelia didelį krovinį.", "Rorres, n.d."),
    ("Hidraulinis principas", " – Paskalio dėsnis: F₁ / A₁ = F₂ / A₂, todėl didesnis stūmoklis duoda didesnę jėgą.", "LibreTexts, n.d."),
], size=21, space=16)

# ── 3. Istorinė raida ───────────────────────────────────────────────────
s = new_slide(3, "Istorinė raida",
              "Raida: nuo Archimedo sverto, antikinių sraigtinių presų, Herono kėlimo mechanizmų ir romėnų kranų per Bramah ir Dudgeon hidrauliką iki žirklinių ir "
              "elektrinių domkratų. Tendencija – nuo jėgos laimėjimo link greičio, saugos ir patogumo.")
table(s, 0.7, 1.6, 11.9, [2.2, 9.7], [
    ["Laikotarpis", "Sprendimas"],
    ["III a. pr. Kr.", "Archimedas aiškina sverto principą („Duokite man atramos tašką…“)"],
    ["I a.", "Graikai ir romėnai naudoja sraigtinius presus alyvuogėms ir vynuogėms spausti"],
    ["~10–75 m.", "Heronas Aleksandrietis „Mechanikoje“ aprašo kėlimo mechanizmus: svertą, suktuvą, skridinį, pleištą ir sraigtą"],
    ["I a. pab.", "Romėnų kranas su mindomu ratu („polyspastos“) kelia iki ~6000 kg"],
    ["1795 m.", "J. Bramah patentuoja hidraulinį presą, veikiantį pagal Paskalio dėsnį"],
    ["1851 m.", "R. Dudgeon patentuoja nešiojamą hidraulinį domkratą"],
    ["1905 m.", "P. J. Harrah patentuoja krumpliastiebinį „Hi-Lift“ domkratą"],
    ["1920 m.", "J. LaFrance pateikia rombo formos (žirklinio) domkrato patento paraišką"],
    ["1949 m.", "W. Watson patobulina žirklinį domkratą – jis tampa sudedamas ir kompaktiškas"],
    ["Šiandien", "Gaminami 12 V elektriniai žirkliniai domkratai"],
], size=15, row_h=0.43)
source_line(s, "Rorres (n.d.); Wikipedia (n.d.-a, n.d.-c, n.d.-d, n.d.-e, n.d.-f); Hi-Lift Jack Company (n.d.-a); AutoIndustriya (n.d.); Pro-Lift-Montagetechnik (n.d.).")

# ── 4. Tipai ir kritinis vertinimas ─────────────────────────────────────
s = new_slide(4, "Domkratų tipai ir jų vertinimas",
              "Palyginome penkis pagrindinius tipus. Pasiteisino savistabdis sraigtas ir hidraulika, o silpnybės – lėtumas, "
              "nestabilumas ir sauga. Gilinsimės į žirklinį, nes jis yra automobilio komplekte.")
table(s, 0.7, 1.6, 11.9, [2.2, 3.4, 3.1, 3.2], [
    ["Tipas", "Veikimo principas", "Pasiteisino", "Trūkumai"],
    ["Sraigtinis", "Sukamas sraigtas kelia stūmoklį", "Savistabdis – krovinys nenusileidžia savaime", "Mažas naudingumo koeficientas (30–40 %), lėtas"],
    ["Krumpliastiebinis", "Svirtis su skląsčiu kyla krumpliastiebiu", "Didelis kėlimo aukštis (iki ~1,5 m)", "Pavojinga rankenos atatranka"],
    ["Žirklinis", "Horizontalus sraigtas suspaudžia rombo formos svirtis", "Lengvas, pigus, telpa automobilyje", "Lėtas, siauras pagrindas, nestabilus"],
    ["Hidraulinis butelinis", "Maža pompa spaudžia alyvą po dideliu stūmokliu", "Didelė keliamoji galia (3–50 t)", "Mažas pagrindas, didelis pradinis aukštis"],
    ["Hidraulinis vežimėlinis", "Hidrocilindras kelia ilgą svirtį", "Greitas ir stabilus (2–4 t)", "Didelis, sunku perkelti"],
], size=15, row_h=0.74, highlight=3)
source_line(s, "Wikipedia (n.d.-b); Joyce/Dayton (n.d.); CJ Pony Parts (n.d.); Hi-Lift Jack Company (n.d.-b); CN Trailer Parts (n.d.); Redline Stands (n.d.).")

# ── 5. Problema ir uždavinys ────────────────────────────────────────────
s = new_slide(5, "Šiuolaikinė problema ir inžinerinis uždavinys",
              "Žirklinis domkratas lėtas ir sunkiai sukamas, ypač pradžioje. Jėga sraigte F = Q / tg θ, "
              "todėl žemoje padėtyje ji kelis kartus didesnė už krovinį. Todėl formuluojame uždavinį.")
tf, st = textbox(s, 0.7, 1.6, 6.0, 0.5, 22, ACCENT, True)
para(tf, st, "Automobilio žirklinio domkrato problemos", first=True)
bullets(s, 0.7, 2.2, 6.1, 4.1, [
    ("Lėtas", " – reikia daug rankenos apsisukimų (mūsų įvertinimu, 60–100)."),
    ("Sunku sukti", " – žemoje padėtyje jėga sraigte kelis kartus didesnė už krovinį (F = Q / tg θ)."),
    ("Nestabilus", " – siauras pagrindas, ypač ant nelygaus ar minkšto paviršiaus."),
    ("Nepatogus", " – sukti tenka nepatogioje padėtyje."),
], size=19, space=12)
rect(s, 7.1, 1.65, 5.5, 4.6, LIGHT)
tf, st = textbox(s, 7.35, 1.85, 5.0, 4.3, 19)
para(tf, st, "Uždavinys", first=True, size=22, bold=True, color=ACCENT)
para(tf, st, "Patobulinti žirklinį domkratą, kad jis 1,5 t automobilį pakeltų greičiau nei per 60 s, "
             "be didelių fizinių pastangų ir stabiliai, išlaikant kompaktiškumą ir saugumą.", space=14)
para(tf, st, "Tikslai:", bold=True, space=4)
for it in ["keliamoji galia ≥ 1500 kg", "kėlimo aukštis 100–400 mm", "kėlimo laikas ≤ 60 s",
           "maitinimas iš 12 V automobilio lizdo", "masė ≤ 4,5 kg"]:
    para(tf, st, it, bullet=True, space=2, size=18)
source_line(s, "CN Trailer Parts (n.d.). Skaičiai ir tikslai – mūsų įvertinimai.")

# ── 6. Sprendimas: elementai ir medžiagos ───────────────────────────────
s = new_slide(6, "Sprendimas: keičiami elementai ir medžiagos",
              "Siūlome elektrinį žirklinį domkratą. Svarbiausi pakeitimai: 12 V variklis, trapecinis sraigtas iš grūdinto plieno "
              "su bronzine veržle, stipresnės svirtys ir platesnis pagrindas. Avarinis rankinis sukimas lieka.")
table(s, 0.7, 1.6, 11.9, [1.9, 2.8, 3.6, 3.6], [
    ["Elementas", "Dabar", "Siūloma", "Kodėl"],
    ["Pavara", "Rankinė rankena", "12 V variklis (~100 W) su reduktoriumi", "Greitai ir be fizinių pastangų"],
    ["Sraigtas", "Paprastas plienas", "Trapecinis sraigtas, plienas 42CrMo4", "Savistabdis; plienas stiprus ir atsparus nuovargiui"],
    ["Veržlė", "Plieninė", "Bronza CuSn12", "Maža trintis, atspari dilimui"],
    ["Svirtys", "Plienas S235 (235 MPa)", "Plienas S420MC (420 MPa)", "Takumo riba ~1,8 karto didesnė"],
    ["Pagrindas", "Siauras", "Platesnis, su guminiu padu", "Stabilesnis, neslysta"],
    ["Sauga", "—", "Galiniai jungikliai, avarinis rankinis sukimas", "Variklis sustoja kraštinėse padėtyse"],
], size=15, row_h=0.64)
source_line(s, "Roton Products (n.d.); Sider Ticino (n.d.); Hengli Automation (n.d.); The World Material (n.d.); Gnee Steel (n.d.); "
               "Mickael (2004). Variklio galia – mūsų skaičiavimas.", y=6.3)

# ── 7. Prieš ir po ──────────────────────────────────────────────────────
s = new_slide(7, "Įprastas ir patobulintas domkratas",
              "Laimime greitį, jėgą ir stabilumą. Kompromisai – didesnė masė ir kaina. Pagal mūsų skaičiavimą variklio galia "
              "apie 100 W, todėl užtenka automobilio 12 V lizdo.")
table(s, 0.7, 1.6, 7.4, [2.6, 2.4, 2.4], [
    ["Parametras", "Įprastas", "Patobulintas"],
    ["Kėlimo laikas", "~2–4 min", "≤ 60 s"],
    ["Rankos jėga", "~100–200 N", "nereikia"],
    ["Pagrindo plotas", "~100 cm²", "~240 cm²"],
    ["Masė", "~3 kg", "~4,5 kg"],
    ["Kaina", "~20–30 €", "~60–80 €"],
], size=17, row_h=0.65)
tf, st = textbox(s, 8.5, 1.6, 4.2, 4.6, 18)
para(tf, st, "Preliminarus skaičiavimas", first=True, size=21, bold=True, color=ACCENT, space=10)
para(tf, st, "Variklio galia:", bold=True, space=2)
para(tf, st, "P ≈ Q·h / (t·η) = 5000·0,25 / (45·0,3) ≈ 93 W", space=12)
para(tf, st, "Srovė iš 12 V lizdo:", bold=True, space=2)
para(tf, st, "I = P / U ≈ 93 / 12 ≈ 8 A", space=12)
para(tf, st, "Savistabdumo sąlyga:", bold=True, space=2)
para(tf, st, "sriegio pakilimo kampas turi būti mažesnis už trinties kampą – tada krovinys nenusileidžia savaime.", space=12)
source_line(s, "Roton Products (n.d.) – savistabdumo sąlyga. Lentelės vertės ir skaičiavimas – mūsų apytiksliai įvertinimai.")

# ── 8. Išvados ──────────────────────────────────────────────────────────
s = new_slide(8, "Išvados",
              "Apibendrinant: domkratai vystėsi nuo jėgos laimėjimo iki saugos ir patogumo. Žirklinį domkratą siūlome "
              "elektrifikuoti ir sustiprinti. Ačiū, laukiame klausimų.")
bullets(s, 0.7, 1.7, 7.4, 4.6, [
    ("Raida", " – nuo jėgos laimėjimo (svertas, sraigtas) per hidrauliką iki elektrinių domkratų."),
    ("Problema", " – žirklinis domkratas kompaktiškas ir savistabdis, bet lėtas, sunkiai sukamas ir nestabilus."),
    ("Sprendimas", " – 12 V elektrinė pavara, trapecinis sraigtas su bronzine veržle, stipresnės svirtys, platesnis pagrindas."),
    ("Kompromisas", " – šiek tiek didesnė masė ir kaina."),
], size=21, space=18)
rect(s, 8.5, 2.6, 4.1, 2.0, LIGHT)
tf, st = textbox(s, 8.5, 2.6, 4.1, 2.0, 36, ACCENT, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
para(tf, st, "Klausimai?", first=True)

# ── 9–10. Šaltiniai (APA) ───────────────────────────────────────────────
REFS = [
    ("AutoIndustriya. (n.d.). ", "Auto essentials: The story of the humble car jack", ". https://www.autoindustriya.com/features/auto-essentials-the-story-of-the-humble-car-jack.html"),
    ("CJ Pony Parts. (n.d.). ", "Types of car jacks", ". https://www.cjponyparts.com/resources/types-of-car-jacks"),
    ("CN Trailer Parts. (n.d.). ", "What are the disadvantages of a scissor jack?", " https://www.cntrailerparts.com/blog/what-are-the-disadvantages-of-a-scissor-jack-372667.html"),
    ("Gnee Steel. (n.d.). ", "EN 10149-2 S420MC hot rolled automotive steel", ". https://www.gneesteel.com/products/automotive-steel/en-10149-2-s420mc-hot-rolled-automotive-steel.html"),
    ("Hengli Automation. (n.d.). ", "What is the material of the nut for a T lead screw?", " https://www.hlt-cnc.com/blog/what-is-the-material-of-the-nut-for-a-t-lead-screw-1318622.html"),
    ("Hi-Lift Jack Company. (n.d.-a). ", "Our history", ". https://hi-lift.com/company/our-history/"),
    ("Hi-Lift Jack Company. (n.d.-b). ", "Hi-Lift jack instructions", " [PDF]. https://hi-lift.com/wp-content/uploads/2016/07/jack_instructions.pdf"),
    ("Joyce/Dayton. (n.d.). ", "Machine screw jacks vs. ball screw jacks", ". https://www.joycedayton.com/blog/machine-screw-jacks-vs-ball-screw-jacks"),
    ("LibreTexts. (n.d.). ", "10.2: Force multiplication", ". https://eng.libretexts.org/Courses/Northeast_Wisconsin_Technical_College/Fluids_1:_Fluid_Power_and_Pneumatics_(NWTC)/10:_Pascal's_Law/10.02:_Force_Multiplication"),
    ("Mickael, E. (2004). ", "Motor driven scissor jack with limit switches", " (U.S. Patent No. 6,695,289). https://patents.google.com/patent/US6695289B1/en"),
    ("Pro-Lift-Montagetechnik. (n.d.). ", "12V electric scissor jack, car jack, 123 mm – 345 mm", ". https://www.pro-lift-montagetechnik.com/12-volt-electric-scissor-jack-car-jack-123mm-345mm-002"),
    ("Redline Stands. (n.d.). ", "Bottle jack vs. floor jack: Differences, pros, and cons", ". https://www.redlinestands.com/blog/bottle-jack-vs-floor-jack/"),
    ("Rorres, C. (n.d.). ", "Quotations about Archimedes' lever", ". New York University. https://math.nyu.edu/Archimedes/Lever/LeverQuotes.html"),
    ("Roton Products. (n.d.). ", "Trapezoidal lead screws: General information", ". https://www.roton.com/products/trapezoidal-lead-screws-nuts/general-information/"),
    ("Sider Ticino. (n.d.). ", "42CrMo4 technical specifications", ". https://siderticino.it/en/steel-datasheets/42crmo4/"),
    ("The World Material. (n.d.). ", "EN 10025-2 S235 steel properties", ". https://www.theworldmaterial.com/en-10025-2-material-s235-steel/"),
    ("Wikipedia. (n.d.-a). ", "Jack (device)", ". https://en.wikipedia.org/wiki/Jack_(device)"),
    ("Wikipedia. (n.d.-b). ", "Jackscrew", ". https://en.wikipedia.org/wiki/Jackscrew"),
    ("Wikipedia. (n.d.-c). ", "Joseph Bramah", ". https://en.wikipedia.org/wiki/Joseph_Bramah"),
    ("Wikipedia. (n.d.-d). ", "Richard Dudgeon", ". https://en.wikipedia.org/wiki/Richard_Dudgeon"),
    ("Wikipedia. (n.d.-e). ", "Simple machine", ". https://en.wikipedia.org/wiki/Simple_machine"),
    ("Wikipedia. (n.d.-f). ", "Treadwheel crane", ". https://en.wikipedia.org/wiki/Treadwheel_crane"),
]
half = (len(REFS) + 1) // 2
for k, chunk in enumerate((REFS[:half], REFS[half:])):
    s = new_slide(9 + k, f"Šaltiniai ({k + 1}/2)")
    tf, st = textbox(s, 0.7, 1.6, 11.9, 5.2, 14)
    for i, (a, title, b) in enumerate(chunk):
        p = para(tf, st, a, first=i == 0, space=7)
        add(p, title, 14, italic=True)
        add(p, b, 14)

prs.save(OUT)
print("saved", OUT)
