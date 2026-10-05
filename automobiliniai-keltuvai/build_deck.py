"""Paprastas PBL TA1 pristatymas „Automobiliniai keltuvai“ (10 skaidrių, be piešinių).

Paleidimas:  python3 build_deck.py   → automobiliniai-keltuvai.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "automobiliniai-keltuvai.pptx"

INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x59, 0x59, 0x59)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
LIGHT = RGBColor(0xDE, 0xEA, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 10


def textbox(slide, x, y, w, h, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.paragraphs[0].alignment = align
    return tf, (size, color, bold)


def para(tf, style, txt, first=False, bullet=False, space=6, size=None, bold=None, color=None):
    s, c, b = style
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space)
    run = p.add_run()
    run.text = ("• " if bullet else "") + txt
    run.font.name, run.font.size = FONT, Pt(size or s)
    run.font.bold = b if bold is None else bold
    run.font.color.rgb = color or c
    return p


def add(p, txt, size, bold=False, italic=False, color=INK):
    r = p.add_run()
    r.text = txt
    r.font.name, r.font.size, r.font.bold, r.font.italic = FONT, Pt(size), bold, italic
    r.font.color.rgb = color
    return r


def bullets(slide, x, y, w, h, items, size=20, space=10):
    """items: (paryškinta dalis, tekstas) arba (paryškinta dalis, tekstas, šaltinis)."""
    tf, st = textbox(slide, x, y, w, h, size)
    for i, it in enumerate(items):
        p = para(tf, st, it[0], first=i == 0, bullet=True, space=space, bold=True)
        add(p, it[1], size)
        if len(it) > 2:
            add(p, f" ({it[2]})", size - 4, color=GREY)
    return tf


def rect(slide, x, y, w, h, color):
    shp = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def new_slide(n, title, notes=""):
    s = prs.slides.add_slide(BLANK)
    tf, st = textbox(s, 0.6, 0.35, 12.1, 0.9, 32, ACCENT, True, anchor=MSO_ANCHOR.BOTTOM)
    para(tf, st, title, first=True)
    rect(s, 0.65, 1.32, 12.0, 0.03, ACCENT)
    tf, st = textbox(s, 11.7, 6.95, 1.1, 0.4, 12, GREY, align=PP_ALIGN.RIGHT)
    para(tf, st, f"{n} / {TOTAL}", first=True)
    tf, st = textbox(s, 0.6, 6.95, 8, 0.4, 12, GREY)
    para(tf, st, "PBL · Automobiliniai keltuvai", first=True)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def source_line(slide, txt, y=6.45):
    tf, st = textbox(slide, 0.7, y, 11.9, 0.5, 13, GREY)
    para(tf, st, "Šaltiniai: " + txt, first=True)


def table(slide, x, y, w, col_w, rows, size=15, row_h=0.5, highlight=None):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shape.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if i == 0 else (LIGHT if i == highlight else WHITE)
            p = cell.text_frame.paragraphs[0]
            add(p, val, size, bold=(i == 0 or j == 0), color=WHITE if i == 0 else INK)
    return tbl


# ── 1. Titulinis ────────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
rect(s, 0, 2.2, 13.333, 2.4, ACCENT)
tf, st = textbox(s, 0.8, 2.35, 11.7, 1.2, 48, WHITE, True, anchor=MSO_ANCHOR.BOTTOM)
para(tf, st, "Automobiliniai keltuvai", first=True)
tf, st = textbox(s, 0.8, 3.55, 11.7, 0.9, 24, WHITE)
para(tf, st, "Raida, problemos ir patobulinimo kryptis", first=True)
tf, st = textbox(s, 0.8, 5.0, 11.7, 1.8, 18, GREY)
para(tf, st, "PBL užduotis · Tarpinis atsiskaitymas Nr. 1", first=True, space=4)
para(tf, st, "Dalykas: Įvadas į specialybę · Mechanikos inžinerija", space=4)
para(tf, st, "Grupė: [Vardas Pavardė, …] · Dėstytojas: G. Viselga · 2026 m.", space=4)
s.notes_slide.notes_text_frame.text = (
    "Laba diena. Mūsų tema – automobiliniai keltuvai, kuriais automobiliai keliami servisuose. Parodysime, kaip keltuvai "
    "vystėsi, kokių problemų jie turi šiandien ir kaip siūlome patobulinti dvikolonį hidraulinį keltuvą.")

# ── 2. Srities pasirinkimas ─────────────────────────────────────────────
s = new_slide(2, "Kodėl pasirinkome automobilinius keltuvus?",
              "Keltuvas yra pagrindinė serviso įranga: jis leidžia pasiekti automobilio apačią. Kadangi automobilis laikomas "
              "virš žmogaus, sauga čia ypač svarbi. Keltuvai veikia pagal Paskalio dėsnį arba sraigto ir veržlės principą.")
bullets(s, 0.7, 1.7, 12, 4.6, [
    ("Pagrindinė serviso įranga", " – dvikolonis keltuvas kelia automobilį už kėbulo, todėl ratai laisvi, o apačia lengvai pasiekiama.", "AIMS Industrial, n.d."),
    ("Didelė atsakomybė", " – automobilis laikomas virš žmogaus; nelaimes dažniausiai lemia operatoriaus klaidos ir apleista priežiūra.", "Automotive Lift Institute, n.d.-b"),
    ("Daug mechanikos vienoje vietoje", " – hidraulinės ir sraigtinės pavaros, metalinės konstrukcijos, fiksatoriai, valdymas."),
    ("Hidraulinis principas", " – Paskalio dėsnis: F₁ / A₁ = F₂ / A₂, todėl didesnis stūmoklis duoda didesnę jėgą.", "LibreTexts, n.d."),
    ("Elektromechaninis principas", " – variklis suka sraigtą kolonoje, o veržlė su vežimėliu kyla aukštyn.", "Automotive Workshop Services, n.d."),
], size=21, space=16)

# ── 3. Istorinė raida ───────────────────────────────────────────────────
s = new_slide(3, "Istorinė raida",
              "Raida: nuo Archimedo sverto, Herono kėlimo mechanizmų ir romėnų kranų per Bramah hidraulinį presą iki Lunati "
              "pirmojo automobilių keltuvo. Vėliau požeminius keltuvus pakeitė paviršiniai, o dabar plinta elektroninis valdymas.")
table(s, 0.7, 1.6, 11.9, [2.2, 9.7], [
    ["Laikotarpis", "Sprendimas"],
    ["III a. pr. Kr.", "Archimedas aiškina sverto principą („Duokite man atramos tašką…“)"],
    ["~10–75 m.", "Heronas Aleksandrietis „Mechanikoje“ aprašo kėlimo mechanizmus: svertą, suktuvą, skridinį, pleištą ir sraigtą"],
    ["I a. pab.", "Romėnų kranas su mindomu ratu („polyspastos“) kelia iki ~6000 kg"],
    ["1795 m.", "J. Bramah patentuoja hidraulinį presą, veikiantį pagal Paskalio dėsnį"],
    ["1925 m.", "P. Lunati, įkvėptas kirpyklos kėdės, sukuria pirmąjį hidraulinį automobilių keltuvą"],
    ["1945 m.", "Įkuriamas Automobilių keltuvų institutas (ALI) – pradedami kurti saugos standartai"],
    ["XX a. 6 deš.", "Atsiranda dvikoloniai keltuvai – pigesnė alternatyva į grindis įleidžiamiems"],
    ["XX a. 8–9 deš.", "Plinta paviršiniai keltuvai; požeminiai kritikuojami dėl alyvos nutekėjimo į gruntą"],
    ["Šiandien", "Mobilios kolonos su belaidžiu sinchronizavimu, elektromechaniniai keltuvai"],
], size=15, row_h=0.47)
source_line(s, "Rorres (n.d.); Wikipedia (n.d.-a, n.d.-b, n.d.-c); Rotary Solutions (n.d.); Automotive Lift Institute (n.d.-a); "
               "Eter Lifts (n.d.); Metro Magazine (n.d.); JMC Automotive Equipment (n.d.).", y=6.4)

# ── 4. Tipai ir kritinis vertinimas ─────────────────────────────────────
s = new_slide(4, "Keltuvų tipai ir jų vertinimas",
              "Palyginome penkis pagrindinius keltuvų tipus. Dvikolonis labiausiai paplitęs, nes leidžia pasiekti visą apačią, "
              "bet jam svarbi abiejų kolonų sinchronizacija ir teisingi kėlimo taškai. Todėl gilinsimės į jį.")
table(s, 0.7, 1.6, 11.9, [2.0, 3.0, 1.4, 2.6, 2.9], [
    ["Tipas", "Veikimo principas", "Galia", "Pasiteisino", "Trūkumai"],
    ["Dvikolonis", "Dvi kolonos, svirtys kelia už kėbulo", "4–8 t", "Ratai laisvi, pasiekiama visa apačia", "Reikia tiksliai parinkti kėlimo taškus"],
    ["Keturkolonis", "Užvažiuojama ant dviejų platformų", "iki ~27 t", "Didelė galia, tinka saugoti automobilius", "Ratai remiasi į platformas"],
    ["Žirklinis", "Žirklinis mechanizmas kelia platformą", "2–3,5 t", "Nereikia aukštų lubų", "Mažas kėlimo aukštis (0,9–1,2 m)"],
    ["Įleidžiamas į grindis", "Stūmokliai įrengti duobėje po grindimis", "2,5–5 t", "Laisvos grindys aplink automobilį", "Alyva gali nutekėti į gruntą"],
    ["Mobilios kolonos", "Atskiros kolonos su akumuliatoriais", "7,5–15 t", "Perkeliamos, tinka sunkiai technikai", "Reikia 4 ir daugiau kolonų"],
], size=15, row_h=0.74, highlight=1)
source_line(s, "AIMS Industrial (n.d.); JMC Automotive Equipment (n.d.); Metro Magazine (n.d.).")

# ── 5. Problema ir uždavinys ────────────────────────────────────────────
s = new_slide(5, "Šiuolaikinė problema ir inžinerinis uždavinys",
              "Daugelyje dvikolonių keltuvų kolonas sinchronizuoja plieniniai lynai. Jie išsitempia ir dyla, todėl vežimėliai "
              "kyla nevienodai, o fiksatoriai užsifiksuoja ne vienu metu. Dažniausia nelaimių priežastis – netinkamai parinkti kėlimo taškai.")
tf, st = textbox(s, 0.7, 1.6, 6.0, 0.5, 22, ACCENT, True)
para(tf, st, "Dvikolonio keltuvo su lynais problemos", first=True)
bullets(s, 0.7, 2.2, 6.1, 4.1, [
    ("Lynai išsitempia ir dyla", " – juos reikia tikrinti kasdien ir reguliuoti kas savaitę ar mėnesį."),
    ("Nevienodas kėlimas", " – vežimėliai kyla nevienodai, fiksatoriai užsifiksuoja ne vienu metu."),
    ("Kėlimo taškai", " – netinkama svorio centro padėtis yra dažniausia dvikolonių keltuvų nelaimių priežastis."),
    ("Svirčių fiksatoriai", " – susidėvėję ar išjungti leidžia svirtims pasislinkti."),
], size=19, space=12)
rect(s, 7.1, 1.65, 5.5, 4.6, LIGHT)
tf, st = textbox(s, 7.35, 1.85, 5.0, 4.3, 19)
para(tf, st, "Uždavinys", first=True, size=22, bold=True, color=ACCENT)
para(tf, st, "Patobulinti dvikolonį hidraulinį keltuvą, kad abu vežimėliai kiltų vienodai be lynų, "
             "fiksatoriai užsifiksuotų kartu, o operatorius būtų įspėtas apie netinkamą automobilio padėtį.", space=14)
para(tf, st, "Tikslai:", bold=True, space=4)
for it in ["keliamoji galia 4 t", "vežimėlių aukščio skirtumas ≤ 5 mm", "kėlimo laikas ≤ 50 s (iki ~1,9 m)",
           "nereikia reguliuoti lynų", "automatiniai mechaniniai fiksatoriai"]:
    para(tf, st, it, bullet=True, space=2, size=18)
source_line(s, "Mohawk Lifts (n.d.); Best Buy Automotive Equipment (n.d.); CarWorship (2026). Tikslai – mūsų.")

# ── 6. Sprendimas: elementai ir medžiagos ───────────────────────────────
s = new_slide(6, "Sprendimas: keičiami elementai ir medžiagos",
              "Siūlome keltuvą su elektronine sinchronizacija: kiekvienoje kolonoje yra padėties jutiklis, o valdiklis reguliuoja "
              "alyvos srautą į cilindrus. Lynų nebelieka. Svirtyse – automatiniai užraktai ir apkrovos jutikliai.")
table(s, 0.7, 1.6, 11.9, [2.0, 2.6, 3.7, 3.6], [
    ["Elementas", "Dabar", "Siūloma", "Kodėl"],
    ["Sinchronizacija", "Plieniniai lynai", "Padėties jutikliai kolonose + valdiklis, reguliuojantis alyvos srautą", "Vežimėliai kyla vienodai, nėra lynų tempimo"],
    ["Fiksatoriai", "Gali užsifiksuoti ne vienu metu", "Automatiniai mechaniniai fiksatoriai su elektromagnetiniu atleidimu", "Užsifiksuoja kartu; EN 1493 reikalauja, kai kėlimas > 500 mm"],
    ["Svirtys", "Užraktai gali susidėvėti", "Automatiniai svirčių užraktai + apkrovos jutikliai atramose", "Svirtys nepasislenka, įspėjama apie netolygią apkrovą"],
    ["Kolonos", "Plieninė konstrukcija", "Plienas S355 (takumo riba 355 MPa)", "Stipri, gerai suvirinama konstrukcija"],
    ["Cilindrų kotai", "—", "Plienas CK45 su kietuoju chromu", "Atsparūs dilimui ir korozijai"],
    ["Valdymas", "Mygtukas", "Laikomas mygtukas (judėjimas tik jį laikant)", "EN 1493 saugos reikalavimas"],
], size=14, row_h=0.66)
source_line(s, "Mohawk Lifts (n.d.); U.S. Patent No. 6,763,916 (n.d.); BSI (2022); CarWorship (2026); Mailhot et al. (n.d.); Steel Available (n.d.); Myler Hydro (n.d.).", y=6.35)

# ── 7. Prieš ir po ──────────────────────────────────────────────────────
s = new_slide(7, "Įprastas ir patobulintas keltuvas",
              "Laimime vienodą kėlimą, mažesnę priežiūrą ir didesnę saugą. Kompromisai – didesnė kaina ir priklausomybė nuo "
              "elektronikos. Pagal mūsų skaičiavimą užtenka 45 mm cilindrų ir apie 2,2 kW siurblio variklio.")
table(s, 0.7, 1.6, 7.4, [2.6, 2.4, 2.4], [
    ["Parametras", "Su lynais", "Patobulintas"],
    ["Sinchronizacija", "Mechaninė (lynai)", "Elektroninė"],
    ["Lynų reguliavimas", "Kas savaitę / mėnesį", "Nereikia"],
    ["Aukščio skirtumas", "Priklauso nuo lynų", "≤ 5 mm (tikslas)"],
    ["Fiksatoriai", "Ne visada kartu", "Užsifiksuoja kartu"],
    ["Kaina", "Mažesnė", "Didesnė"],
    ["Gedimo rizika", "Lynų nusidėvėjimas", "Jutiklių gedimas"],
], size=16, row_h=0.6)
tf, st = textbox(s, 8.5, 1.6, 4.2, 4.6, 17)
para(tf, st, "Preliminarus skaičiavimas", first=True, size=21, bold=True, color=ACCENT, space=10)
para(tf, st, "Jėga vienam cilindrui:", bold=True, space=2)
para(tf, st, "F = 4000 · 9,81 / 2 ≈ 19,6 kN", space=10)
para(tf, st, "Cilindro skersmuo (p = 150 bar):", bold=True, space=2)
para(tf, st, "D = √(4F / (π·p)) ≈ 41 mm → renkamės 45 mm", space=10)
para(tf, st, "Siurblio variklio galia:", bold=True, space=2)
para(tf, st, "P = F·v / η = 39 240 · 0,04 / 0,7 ≈ 2,2 kW", space=10)
source_line(s, "Mohawk Lifts (n.d.). Lentelės vertės ir skaičiavimas – mūsų apytiksliai įvertinimai.")

# ── 8. Išvados ──────────────────────────────────────────────────────────
s = new_slide(8, "Išvados",
              "Apibendrinant: keltuvai vystėsi nuo požeminių hidraulinių iki paviršinių ir elektroniškai valdomų. "
              "Dvikolonį keltuvą siūlome patobulinti elektronine sinchronizacija ir automatiniais fiksatoriais. Ačiū, laukiame klausimų.")
bullets(s, 0.7, 1.7, 7.4, 4.6, [
    ("Raida", " – nuo pirmojo hidraulinio keltuvo (1925 m.) per paviršinius dvikolonius iki elektroninio valdymo."),
    ("Problema", " – lynai išsitempia, vežimėliai kyla nevienodai, o dažniausia nelaimių priežastis – netinkami kėlimo taškai."),
    ("Sprendimas", " – elektroninė sinchronizacija, automatiniai fiksatoriai ir svirčių užraktai, apkrovos jutikliai."),
    ("Kompromisas", " – didesnė kaina ir priklausomybė nuo elektronikos."),
], size=21, space=18)
rect(s, 8.5, 2.6, 4.1, 2.0, LIGHT)
tf, st = textbox(s, 8.5, 2.6, 4.1, 2.0, 36, ACCENT, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
para(tf, st, "Klausimai?", first=True)

# ── 9–10. Šaltiniai (APA) ───────────────────────────────────────────────
REFS = [
    ("AIMS Industrial. (n.d.). ", "Vehicle hoist guide: 2-post, 4-post, scissor & car lifts", ". https://aimsindustrial.com.au/blogs/product-guides/vehicle-hoist-guide"),
    ("Automotive Lift Institute. (n.d.-a). ", "ALI history timeline", ". https://www.autolift.org/automotive-lift-institute-history/"),
    ("Automotive Lift Institute. (n.d.-b). ", "Lift safety", ". https://www.autolift.org/news/lift-safety/"),
    ("Automotive Workshop Services. (n.d.). ", "A major overhaul on a two-post electro-mechanical vehicle hoist", ". https://automotiveworkshopservices.com.au/vehicle-and-truck-hoists/a-major-overhaul-on-a-two-post-electro-mechanical-vehicle-hoist/"),
    ("Best Buy Automotive Equipment. (n.d.). ", "Which is safer: 2-post or 4-post lift?", " https://www.bestbuyautoequipment.com/which-is-safer-2-post-or-4-post-lift-safety-comparison/"),
    ("BSI. (2022). ", "Vehicle lifts", " (BS EN 1493:2022). https://www.en-standard.eu/bs-en-1493-2022-vehicle-lifts/"),
    ("CarWorship. (2026). ", "How vehicle lift safety locks prevent accidents during automotive maintenance", ". https://www.carworship.com/tips/how-vehicle-lift-safety-locks-prevent-catastrophic-failures-during-maintenance/"),
    ("Eter Lifts. (n.d.). ", "Origin of the two-post car lift", ". https://www.eterlifts.com/news/origin-of-the-two-post-car-lift-84982494.html"),
    ("JMC Automotive Equipment. (n.d.). ", "8 types of car lifts compared", ". https://jmcautomotiveequipment.com/pages/jmc-equipment-buyers-guide/vehicle-lifts-different-types-and-how-to-choose-the-right-one.html"),
    ("LibreTexts. (n.d.). ", "10.2: Force multiplication", ". https://eng.libretexts.org/Courses/Northeast_Wisconsin_Technical_College/Fluids_1:_Fluid_Power_and_Pneumatics_(NWTC)/10:_Pascal's_Law/10.02:_Force_Multiplication"),
    ("Mailhot, L., Mitchell, N., & Jean, S. (n.d.). ", "Arm restraints for vehicle lift and vehicle lift including the same", " (U.S. Patent No. 11,613,452). https://patents.google.com/patent/US11613452"),
    ("Metro Magazine. (n.d.). ", "Diverse vehicle lift designs fit variety of needs", ". https://www.metro-magazine.com/articles/diverse-vehicle-lift-designs-fit-variety-of-needs"),
    ("", "Method and apparatus for synchronizing a vehicle lift", " (U.S. Patent No. 6,763,916). (n.d.). https://image-ppubs.uspto.gov/dirsearch-public/print/downloadPdf/6763916"),
    ("Mohawk Lifts. (n.d.). ", "Two post lifts: Hydraulic vs. cable synchronization", ". https://mohawklifts.com/blog/two-post-lifts-hydraulic-vs-cable-synchronization/"),
    ("Myler Hydro. (n.d.). ", "CK45 hard chrome-plated rod", ". https://mylerhydro.com/en/ck45-chrome-rod/"),
    ("Rorres, C. (n.d.). ", "Quotations about Archimedes' lever", ". New York University. https://math.nyu.edu/Archimedes/Lever/LeverQuotes.html"),
    ("Rotary Solutions. (n.d.). ", "About Rotary Solutions", ". https://rotarysolutions.com/about/"),
    ("Steel Available. (n.d.). ", "Steel S355: The ideal material for heavy-duty applications", ". https://www.steelavailable.com/en/steel-s355-the-ideal-material-for-heavy-duty-applications/"),
    ("Wikipedia. (n.d.-a). ", "Joseph Bramah", ". https://en.wikipedia.org/wiki/Joseph_Bramah"),
    ("Wikipedia. (n.d.-b). ", "Simple machine", ". https://en.wikipedia.org/wiki/Simple_machine"),
    ("Wikipedia. (n.d.-c). ", "Treadwheel crane", ". https://en.wikipedia.org/wiki/Treadwheel_crane"),
]
half = (len(REFS) + 1) // 2
for k, chunk in enumerate((REFS[:half], REFS[half:])):
    s = new_slide(9 + k, f"Šaltiniai ({k + 1}/2)")
    tf, st = textbox(s, 0.7, 1.6, 11.9, 5.2, 13)
    for i, (a, title, b) in enumerate(chunk):
        p = para(tf, st, a, first=i == 0, space=6)
        add(p, title, 13, italic=True)
        add(p, b, 13)

prs.save(OUT)
print("saved", OUT)
