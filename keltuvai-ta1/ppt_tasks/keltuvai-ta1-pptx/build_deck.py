"""Build the editable PPTX for 'Keltuvai: raida ir tobulinimo galimybės' (TA1).

Reproducible Build Mode script (ppt-design-skill + pptx-designer public API).
Design direction: Cobalt Grid (theme-lock.yaml v1). Coordinates are authored on
the same 1920×1080 grid as the HTML deck and converted with px() → inches.

Run:  python build_deck.py   → ../../keltuvai-ta1.pptx
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from PIL import Image, ImageDraw
from pptx.util import Pt
from pptx_designer import Presentation, validate_resolved_theme
from pptx_designer.renderer.theme import ThemeComposer
from pptx_designer.tools.images import cover_image
from pptx_designer.tools.shapes import add_slide, oval, rect, triangle
from pptx_designer.tools.text import text

HERE = Path(__file__).resolve().parent
OUT = HERE.parents[1] / "keltuvai-ta1.pptx"

# ── Locked tokens (theme-lock.yaml v1) ─────────────────────────────────
PAPER, INK, INK_SOFT = "#F0EBDE", "#1F2BE0", "#5560E5"
INK_FAINT = "#C8C6DE"      # 18 % cobalt on paper, pre-blended
TINT = "#DFDCE0"           # 8 % cobalt on paper — chosen-column fill
SERIF, SANS, MONO = "Georgia", "Calibri", "Consolas"
TOTAL = 14


def px(v: float) -> float:
    """1920-px stage units → inches (13.333 in wide)."""
    return v / 144.0


def resolved_theme() -> dict:
    """Compose once, pin the locked tokens, validate, and save with a fingerprint."""
    t = ThemeComposer().compose(layout="standard", seed=17)
    t["colors"].update({"primary": INK, "on-primary": PAPER, "secondary": INK_SOFT, "accent": INK,
                        "background": PAPER, "foreground": INK, "muted": TINT,
                        "muted-foreground": INK_SOFT, "border": INK_FAINT})
    t["typography"] = {"heading": SERIF, "body": SANS}
    t.setdefault("semantic_roles", {}).update({"background": PAPER, "surface": PAPER, "ink": INK,
                                               "muted": INK_SOFT, "accent": INK, "border": INK_FAINT,
                                               "data-series-1": INK, "data-series-2": INK_SOFT})
    validate_resolved_theme(t)
    blob = json.dumps(t, sort_keys=True, default=str)
    (HERE / "resolved-theme-v1.json").write_text(json.dumps(t, indent=2, default=str), encoding="utf-8")
    print("resolved theme sha256:", hashlib.sha256(blob.encode()).hexdigest())
    return t


def grid_background() -> Path:
    """Cream paper with a 48-px cobalt graph-paper grid (decorative background only)."""
    path = HERE / "assets" / "grid-paper.png"
    path.parent.mkdir(exist_ok=True)
    img = Image.new("RGB", (3840, 2160), PAPER)
    d = ImageDraw.Draw(img)
    line = "#DBD8DE"  # 10 % cobalt on paper
    for x in range(0, 3840, 96):
        d.line([(x, 0), (x, 2160)], fill=line, width=2)
    for y in range(0, 2160, 96):
        d.line([(0, y), (3840, y)], fill=line, width=2)
    img.save(path, optimize=True)
    return path


# ── Drawing helpers (all built on public pptx_designer helpers) ───────────
def t(s, x, y, w, h, txt, size, font=SANS, color=INK, align="left", italic=False, bold=False,
      spacing=None, anchor="top", line=None):
    """Text box at stage px; size in stage px (pt = px / 2). Zero insets for exact placement."""
    box = text(s, px(x), px(y), px(w), px(h), txt, font_size=size / 2, color=color, bold=bold,
               align=align, font_name=font, anchor=anchor)
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    run = tf.paragraphs[0].runs[0]
    run.font.italic = italic
    if spacing is not None:
        run.font._rPr.set("spc", str(int(spacing * size / 2 * 100)))  # em → 1/100 pt
    if line:
        tf.paragraphs[0].line_spacing = line
    return box


def add_run(box, txt, size, font=SANS, color=INK, italic=False, bold=False, new_par=False, line=None,
            space_before=0):
    """Append a run (or a new paragraph) to a text box created by t()."""
    tf = box.text_frame
    p = tf.add_paragraph() if new_par else tf.paragraphs[-1]
    if new_par:
        p.alignment = tf.paragraphs[0].alignment
        p.space_before = Pt(space_before)
        if line:
            p.line_spacing = line
    r = p.add_run()
    r.text = txt
    r.font.size, r.font.name, r.font.italic, r.font.bold = Pt(size / 2), font, italic, bold
    r.font.color.rgb = __import__("pptx").dml.color.RGBColor.from_string(color.lstrip("#"))
    return r


def hline(s, x, y, w, weight=1.5, color=INK):
    return rect(s, px(x), px(y), px(w), px(weight), color)


def vline(s, x, y, h, weight=1.5, color=INK):
    return rect(s, px(x), px(y), px(weight), px(h), color)


def box_outline(s, x, y, w, h, weight=3, fill=PAPER, color=INK):
    shp = rect(s, px(x), px(y), px(w), px(h), fill, line=color)
    shp.line.width = Pt(weight / 2)
    return shp


def ring(s, cx, cy, r, weight=3, color=INK):
    shp = oval(s, px(cx - r), px(cy - r), px(2 * r), px(2 * r), PAPER, line=color)
    shp.line.width = Pt(weight / 2)
    return shp


def dot(s, cx, cy, r=3, color=INK):
    return oval(s, px(cx - r), px(cy - r), px(2 * r), px(2 * r), color)


def new_slide(prs, bg):
    s = add_slide(prs)
    cover_image(s, 0, 0, 13.333, 7.5, str(bg))
    hline(s, 80, 44, 1760)          # top hairline
    hline(s, 80, 1038, 1760)        # bottom hairline
    return s


def chrome(s, n, src=None):
    t(s, 1540, 994, 300, 24, f"{n:02d} / {TOTAL}", 18, MONO, align="right", spacing=0.06)
    if src:
        t(s, 80, 986, 1420, 44, src, 20, MONO, spacing=0.02, line=1.0)


def topbar(s, title, lab):
    t(s, 80, 86, 1400, 100, title, 76, SERIF, spacing=-0.008)
    t(s, 1340, 158, 500, 24, lab.upper(), 17, SANS, align="right", bold=True, spacing=0.18)
    hline(s, 80, 196, 1760)


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


def glitch(s, x0, top, height, steps, from_left=False):
    """Stair-stepped scanline column (Cobalt Grid signature)."""
    widths = [4, 4, 10, 4, 18, 4, 4, 26, 4, 10, 40, 4, 8, 56, 4, 4, 72, 12, 4, 90]
    x = 0
    for i in range(steps):
        w = widths[i % len(widths)]
        y = top + round(i * height / (steps + 4))
        left = x0 + x if from_left else x0 - x - w
        rect(s, px(left), px(y), px(w), px(top + height - y), INK)
        x += w + 8


def qr(s, x, y, cell=13):
    pat = "1110011110100101101101011000110101011010011011011001011111001101"
    rect(s, px(x - 4), px(y - 4), px(8 * (cell + 1.5) + 6), px(8 * (cell + 1.5) + 6), PAPER)
    for i, c in enumerate(pat):
        if c == "1":
            rect(s, px(x + (i % 8) * (cell + 1.5)), px(y + (i // 8) * (cell + 1.5)), px(cell), px(cell), INK)


# ── Slides ────────────────────────────────────────────────────────────
def s01_cover(prs, bg):
    s = new_slide(prs, bg)
    for x, txt, al in ((80, "ĮVADAS Į SPECIALYBĘ", "left"), (660, "MECHANIKOS INŽINERIJA", "center"),
                       (1340, "TA1 · 2026", "right")):
        t(s, x, 66, 500, 24, txt, 19, MONO, align=al, spacing=0.08)
    t(s, 80, 196, 1000, 24, "PBL UŽDUOTIS · TARPINIS ATSISKAITYMAS NR. 1", 18, SANS, bold=True, spacing=0.18)
    title = t(s, 74, 236, 1180, 440, "Keltuvai: raida ir ", 136, SERIF, spacing=-0.02, line=0.86)
    add_run(title, "tobulinimo", 136, SERIF, INK_SOFT, italic=True)
    add_run(title, " galimybės", 136, SERIF)
    t(s, 80, 736, 820, 100, "Nuo Koliziejaus kabestanų iki traukos liftų — ir kaip šiandien sumažinti judančią masę.",
      32, SANS, line=1.25)
    t(s, 80, 860, 360, 30, "[Grupės pavadinimas]", 19, MONO)
    t(s, 420, 860, 480, 30, "[Vardas Pavardė, Vardas Pavardė, …]", 19, MONO, align="right")
    hline(s, 80, 900, 820, 1, INK_FAINT)
    glitch(s, 1840, 110, 820, 20)
    qr(s, 1330, 180)
    chrome(s, 1)
    notes(s, "Sveiki. Mūsų tema — keltuvai: kaip jie vystėsi ir kur šiandien galima juos patobulinti. "
             "Pristatysime TA1 dalį: srities pasirinkimą, raidos analizę, problemą ir inžinerinio sprendimo pagrindimą.")


def s02_why(prs, bg):
    s = new_slide(prs, bg)
    topbar(s, "Kodėl pasirinkome keltuvus", "PBL 1 · Sritis")
    b = t(s, 80, 236, 840, 200, "Keltuvai naudojami pastatuose, automobilių servisuose, statybvietėse ir sandėliuose. "
          "Visų jų uždavinys tas pats: ", 32, SANS, line=1.3)
    add_run(b, "saugiai pakelti krovinį ar žmones ir sunaudoti kuo mažiau energijos.", 32, SANS, bold=True)
    t(s, 1000, 236, 840, 200, "Šioje srityje susitinka pavaros, medžiagos, sauga ir energetika, todėl ji tinka "
      "pradedantiems mechanikos inžinieriams.", 32, SANS, line=1.3)
    hline(s, 80, 560, 1760)
    stats = [("3–8", " %", "", "visos pastato elektros energijos sunaudoja liftai ir eskalatoriai"),
             ("18", " TWh", "", "elektros energijos per metus sunaudoja Europos Sąjungos liftai"),
             ("66", " %", "iki ", "šios energijos būtų galima sutaupyti taikant geriausias technologijas")]
    for i, (num, unit, pre, cap) in enumerate(stats):
        x = 80 + i * 587
        if i:
            vline(s, x - 30, 560, 390, 1, INK_FAINT)
        n = t(s, x, 590, 560, 200, pre, 72, SERIF) if pre else t(s, x, 590, 560, 200, num, 180, SERIF, spacing=-0.02)
        if pre:
            add_run(n, num, 180, SERIF)
        add_run(n, unit, 72, SERIF)
        t(s, x, 848, 520, 100, cap, 30, SANS, line=1.15)
    chrome(s, 2, "Šaltiniai: de Almeida ir kt. (2010); Krakowski ir Ruta (2018)")
    notes(s, "Keltuvai yra visur, o jų uždavinys visada tas pats — saugiai pakelti ir sunaudoti kuo mažiau energijos. "
             "Liftai sunaudoja 3–8 % pastato elektros, ES — apie 18 TWh per metus, o geriausios technologijos leistų sutaupyti iki 66 %.")


def s03_history(prs, bg):
    s = new_slide(prs, bg)
    topbar(s, "Istorinė raida", "PBL 2 · Raidos seka")
    cols = (80, 250, 810, 1610)
    for x, h, al, w in zip(cols, ("METAI", "SPRENDIMAS", "KAS PASIKEITĖ", "KRYPTIS"),
                           ("left", "left", "left", "right"), (160, 540, 780, 230)):
        t(s, x, 232, w, 20, h, 15, SANS, bold=True, spacing=0.18, align=al)
    hline(s, 80, 262, 1760)
    rows = [("I a.", "Koliziejaus keltuvai", "Vergai sukdavo kabestanus; lynai ir atsvarai kėlė narvus į areną", "raumenų jėga", 0),
            ("1846", "Hidraulinis kranas (W. Armstrongas)", "Hidraulika imta naudoti ir keltuvams", "hidraulika", 0),
            ("1854", "Saugos stabdys (E. Otisas)", "Nutrūkus lynui, platformą sulaiko spyruokliniai fiksatoriai", "→ sauga", 1),
            ("1877", "Trinties skriemulys (F. Koepe)", "Būgną pakeitė skriemulys, kuris lyną varo trintimi", "→ trauka", 1),
            ("1880", "Elektrinis keltuvas (W. von Siemensas)", "Pirmasis elektra varomas keltuvas", "elektra", 0),
            ("1902", "Bereduktorinis traukos liftas (Otis)", "Aukštuose pastatuose išstūmė hidraulinį", "→ elektra", 1),
            ("1925", "Automobilių keltuvas (P. Lunati)", "Užpatentuotas pirmasis hidraulinis automobilių keltuvas", "hidraulika", 0),
            ("1996", "Liftas be mašinų patalpos (KONE)", "Plokščias bereduktorinis variklis įrengtas pačioje šachtoje", "kompaktiškumas", 0),
            ("2013", "Anglies pluošto lynas (KONE)", "Liftas gali kilti iki 1 km — dvigubai aukščiau nei anksčiau", "lengvinimas", 0)]
    y = 266
    for yr, nm, ch, tg, hl in rows:
        if hl:
            rect(s, px(80), px(y), px(1760), px(62), TINT)
        t(s, 80, y + 12, 160, 44, yr, 38, SERIF)
        t(s, 250, y + 16, 540, 40, nm, 30, SERIF)
        t(s, 810, y + 20, 780, 34, ch, 24, SANS)
        t(s, 1610, y + 24, 230, 24, tg, 18, MONO, align="right")
        y += 66
        hline(s, 80, y - 3, 1760, 1, INK_FAINT)
    chrome(s, 3, "Šaltiniai: Klein (2015); ERIH (b. d.); Otis (1861); The Elevator Museum (b. d.); ETHW (2015); "
                 "Rotary Solutions (b. d.); KONE Corporation (2026); FMLink (2013)")
    notes(s, "Raida: nuo raumenų jėgos Koliziejuje iki hidraulikos, o lemiami lūžiai — Otiso saugos stabdys 1854 m., "
             "Koepe trinties skriemulys 1877 m. ir elektrinė traukos pavara. Naujausia kryptis — kompaktiškumas ir lengvesni elementai.")


def s04_schemes(prs, bg):
    s = new_slide(prs, bg)
    topbar(s, "Veikimo schemos ir vertinimas", "PBL 2 · Kritinis vertinimas")
    col_w, gap = 410, 40
    base_y = 250
    specs = [
        ("Raumenų jėga\n(kabestanas, gervė)", "✗ ATSISAKYTA", False,
         "Paprasta, nereikia kuro ar elektros", "Lėta, vienam keltuvui reikia kelių žmonių"),
        ("Būgninis\nkeltuvas", "✗ ATSISAKYTA", False,
         "Paprasta konstrukcija", "Didėjant aukščiui būgnas tampa milžiniškas; didelė lyno trūkimo rizika"),
        ("Hidraulinis\nkeltuvas", "≈ RIBOTAI", False,
         "Didelė jėga", "Grunto taršos ir gaisro rizika; aukštuose pastatuose jį išstūmė elektrinis liftas"),
        ("Trauka su atsvaru\nir saugos stabdžiu", "✓ PASITEISINO", True,
         "Tapo šiuolaikinės liftų inžinerijos pagrindu", "Sunkūs plieniniai lynai; skriemulys turi būti ≥ 40 lyno skersmenų"),
    ]
    for i, (name, verdict, ok, plus, minus) in enumerate(specs):
        x = 80 + i * (col_w + gap + 10)
        if i:
            vline(s, x - 25, 225, 700, 1, INK_FAINT)
        draw_scheme(s, i, x, base_y, col_w)
        nm = t(s, x, 530, col_w, 100, name.split("\n")[0], 36, SERIF, line=1.0)
        add_run(nm, name.split("\n")[1], 36, SERIF, new_par=True, line=1.0)
        vw = 36 + len(verdict) * 13
        if ok:
            rect(s, px(x), px(640), px(vw), px(34), INK)
        else:
            box_outline(s, x, 640, vw, 34, weight=3)
        t(s, x + 12, 647, vw - 12, 22, verdict, 17, MONO, PAPER if ok else INK, spacing=0.02)
        t(s, x, 700, 30, 30, "+", 24, MONO)
        t(s, x + 30, 698, col_w - 30, 70, plus, 24, SANS, line=1.1)
        t(s, x, 790, 30, 30, "−", 24, MONO)
        t(s, x + 30, 788, col_w - 30, 120, minus, 24, SANS, line=1.1)
    chrome(s, 4, "Principinės schemos supaprastintos. Vertinimas pagal Klein (2015); The Elevator Museum (b. d.); "
                 "KONE Corporation (b. d., 2026); ETHW (2015); Kalliomaki (2022); Muñoz (2020)")
    notes(s, "Keturios principinės schemos. Raumenų jėga — lėta. Būgnas — aukštį riboja jo dydis. Hidraulika — didelė jėga, "
             "bet taršos ir gaisro rizika, aukštuose pastatuose išstumta. Pasiteisino trauka su atsvaru ir saugos stabdžiu, "
             "tačiau jos trūkumas — sunkūs plieniniai lynai ir didelis skriemulys.")


def draw_scheme(s, kind, x, y, w):
    """Principal schematics, 400×250 px boxes, cobalt line work."""
    L = lambda a, b, c, weight=3: hline(s, x + a, y + b, c, weight)   # noqa: E731
    V = lambda a, b, c, weight=3: vline(s, x + a, y + b, c, weight)   # noqa: E731
    lab = lambda a, b, txt: t(s, x + a, y + b, 120, 20, txt, 15, MONO)  # noqa: E731
    L(20, 236, 360)  # ground
    if kind == 0:    # capstan → rope over head pulley → cage
        box_outline(s, x + 60, y + 150, 46, 86)
        L(10, 176, 146)
        L(106, 196, 110)
        V(214, 36, 162)
        L(214, 36, 70)
        ring(s, x + 300, y + 36, 22)
        dot(s, x + 300, y + 36)
        V(322, 36, 84)
        box_outline(s, x + 276, y + 120, 92, 84)
        L(276, 162, 92, 1.5)
        V(322, 120, 84, 1.5)
        L(240, 8, 120)
    elif kind == 1:  # drum hoist driven by motor M
        L(60, 30, 280)
        box_outline(s, x + 150, y + 36, 100, 56)
        for dy in (50, 64, 78):
            L(150, dy, 100, 1.5)
        L(250, 64, 40)
        box_outline(s, x + 290, y + 46, 44, 36)
        lab(304, 88, "M")
        V(200, 92, 78)
        box_outline(s, x + 130, y + 170, 140, 16)
    elif kind == 2:  # hydraulic: pump → cylinder → piston → platform
        box_outline(s, x + 120, y + 40, 160, 16)
        V(196, 56, 74, 8)
        box_outline(s, x + 176, y + 130, 48, 106)
        L(224, 214, 76)
        V(300, 196, 18)
        box_outline(s, x + 286, y + 168, 60, 28)
        lab(282, 146, "siurblys")
    else:            # traction sheave, cabin + counterweight, safety gear on rail
        L(60, 10, 280)
        ring(s, x + 200, y + 44, 34)
        dot(s, x + 200, y + 44)
        V(166, 44, 66)
        V(233, 44, 86)
        box_outline(s, x + 112, y + 110, 108, 92)
        rect(s, px(x + 222), px(y + 130), px(28), px(72), INK)
        V(96, 60, 176, 1.5)
        triangle(s, px(x + 97), px(y + 184), px(16), px(22), INK)
        lab(262, 160, "atsvaras")
        lab(0, 214, "stabdys")


def s05_verdict(prs, bg):
    s = new_slide(prs, bg)
    topbar(s, "Kas pasiteisino, o kas — ne", "PBL 2 · Išvada")
    cols = [(80, "✗", "NETINKAMI SPRENDIMAI",
             [("Raumenų jėga", "", "Per lėta ir per daug darbo jėgos reikalaujanti"),
              ("Būgnas lynui vynioti", "", "Kėlimo aukštį riboja būgno dydis"),
              ("Hidraulika aukštuose pastatuose", "", "Išstumta elektrinio traukos lifto")]),
            (996, "✓", "LABIAUSIAI PASITEISINO",
             [("Automatinis saugos stabdys", "1854", "Nutrūkus lynui kabina nebekrenta"),
              ("Trinties skriemulys su atsvaru", "1877", "Pakeitė būgną — aukštis nebėra ribojamas"),
              ("Elektrinė traukos pavara", "1902", "Aukštuose pastatuose išstūmė hidraulinę")])]
    for x, mark, head, items in cols:
        t(s, x, 228, 60, 60, mark, 56, SERIF)
        t(s, x + 64, 252, 600, 24, head, 19, SANS, bold=True, spacing=0.18)
        hline(s, x, 296, 844)
        y = 300
        for h4, yr, p in items:
            hb = t(s, x, y + 26, 844, 50, h4, 42, SERIF)
            if yr:
                add_run(hb, "   " + yr, 18, MONO)
            t(s, x, y + 82, 844, 36, p, 27, SANS)
            y += 142
            hline(s, x, y, 844, 1, INK_FAINT)
    hline(s, 80, 850, 1760)
    t(s, 80, 866, 80, 70, "→", 56, SERIF)
    br = t(s, 160, 878, 1680, 60, "Likęs trūkumas — ", 34, SERIF)
    add_run(br, "sunkūs plieniniai lynai ir didelis skriemulys.", 34, SERIF, italic=True)
    add_run(br, " Čia prasideda mūsų problema.", 34, SERIF)
    chrome(s, 5, "Šaltiniai: Klein (2015); The Elevator Museum (b. d.); KONE Corporation (b. d., 2026); ETHW (2015); "
                 "Kalliomaki (2022); Muñoz (2020)")
    notes(s, "Apibendrinant vertinimą: atsisakyta raumenų jėgos, būgno ir hidraulikos aukštuose pastatuose. "
             "Pasiteisino saugos stabdys, trinties skriemulys su atsvaru ir elektrinė traukos pavara. "
             "Liko vienas trūkumas — sunkūs lynai ir didelis skriemulys. Nuo čia prasideda mūsų problema.")


def s06_problem(prs, bg):
    s = new_slide(prs, bg)
    topbar(s, "Problema: per didelė judanti masė", "PBL 1 · Problema")
    t(s, 80, 230, 960, 40, "Kiekvieno važiavimo metu variklis turi pajudinti ne tik krovinį:", 28, SANS)
    box_outline(s, 80, 292, 960, 238)
    t(s, 118, 322, 890, 120, "Judanti masė = kabina + krovinys + atsvaras + lynai", 46, SERIF, line=1.1)
    t(s, 118, 460, 890, 40, "Atsvaras ≈ kabina + (0,4–0,5) × vardinis krovinys", 26, MONO)
    t(s, 80, 566, 960, 40, "Atsvaras priklauso nuo kabinos masės, todėl kabinos palengvinimas sumažina ir atsvarą:", 24, SANS)
    parts = [(80, "−100", "KG KABINA", False), (410, "−100", "KG ATSVARAS", False), (748, "−200", "KG JUDANČIOS MASĖS", True)]
    for x, v, lbl, res in parts:
        if res:
            rect(s, px(x), px(652), px(292), px(116), INK)
        t(s, x + (14 if res else 0), 658, 290, 110, v, 100, SERIF, PAPER if res else INK)
        t(s, x, 786, 320, 24, lbl, 18, SANS, bold=True, spacing=0.16)
    t(s, 350, 690, 50, 60, "+", 56, SERIF, align="center")
    t(s, 684, 690, 50, 60, "=", 56, SERIF, align="center")
    vline(s, 1120, 222, 740)
    t(s, 1165, 224, 600, 24, "DIDELĖ MASĖ LEMIA", 18, SANS, bold=True, spacing=0.18)
    y = 270
    for i, item in enumerate(["galingesnį ir brangesnį variklį", "didesnes energijos sąnaudas įsibėgėjant ir stabdant",
                              "storesnius lynus ir didesnį skriemulį", "didesnes stabdžių ir konstrukcijos apkrovas"]):
        t(s, 1165, y + 30, 60, 30, f"0{i + 1}", 20, MONO)
        t(s, 1229, y + 24, 610, 90, item, 31, SANS, line=1.15)
        y += 118
        hline(s, 1165, y, 675, 1, INK_FAINT)
    chrome(s, 6, "Šaltiniai: TK Elevator (b. d.); Cloux ir kt. (1999); KONE Corporation (b. d.)")
    notes(s, "Problema — judanti masė. Variklis kiekvieną kartą judina kabiną, krovinį, atsvarą ir lynus. "
             "Kadangi atsvaras parenkamas pagal kabinos masę, palengvinę kabiną 100 kg, sumažiname ir atsvarą — iš viso 200 kg.")


def s07_task(prs, bg):
    s = new_slide(prs, bg)
    topbar(s, "Inžinerinis uždavinys", "PBL 1–2 · Uždavinys")
    task = t(s, 80, 226, 1700, 220, "Sumažinti keltuvo judančių dalių masę, pakeičiant ", 58, SERIF, line=1.05)
    add_run(task, "rėmo medžiagą", 58, SERIF, INK_SOFT, italic=True)
    add_run(task, " ir ", 58, SERIF)
    add_run(task, "kėlimo elementą", 58, SERIF, INK_SOFT, italic=True)
    add_run(task, ", išlaikant tą pačią keliamąją galią ir saugos atsargos koeficientus.", 58, SERIF)
    cols = (80, 150, 640, 1250)
    for x, h, w in zip(cols, ("NR.", "KEIČIAMAS ELEMENTAS", "DABAR → SIŪLOMA", "KEIČIAMI PARAMETRAI"), (80, 420, 590, 600)):
        t(s, x, 486, w, 20, h, 15, SANS, bold=True, spacing=0.18)
    hline(s, 80, 516, 1760)
    rows = [("01", "Kabinos ar platformos rėmas", "Konstrukcinis plienas", "aliuminio lydinys",
             "Profilių matmenys, sienelės storis, jungčių tipas"),
            ("02", "Kėlimo elementas", "Plieninis lynas", "dengtas plieninis diržas",
             "Diržo plotis, diržų skaičius, atsargos koeficientas"),
            ("03", "Skriemulys ir variklis", "Didelis skriemulys", "mažesnis skriemulys",
             "Skriemulio skersmuo, variklio sukimo momentas")]
    y = 520
    for n, el, old, new, prm in rows:
        t(s, 80, y + 34, 80, 24, n, 20, MONO)
        t(s, 150, y + 28, 480, 50, el, 34, SERIF)
        ft = t(s, 640, y + 32, 600, 40, old, 26, SANS, INK_SOFT)
        add_run(ft, "  →  ", 26, SANS)
        add_run(ft, new, 26, SANS, bold=True)
        t(s, 1250, y + 34, 590, 40, prm, 24, SANS)
        y += 102
        hline(s, 80, y, 1760, 1, INK_FAINT)
    t(s, 80, 880, 1700, 80, "Plieniniam lynui skriemulio skersmuo turi būti bent 40 kartų didesnis už lyno skersmenį. "
      "Plonas diržas leidžia naudoti daug mažesnį skriemulį, todėl varikliui reikia mažesnio sukimo momento.", 25, SANS, line=1.2)
    chrome(s, 7, "Šaltinis: Muñoz (2020)")
    notes(s, "Mūsų uždavinys — sumažinti judančią masę keičiant rėmo medžiagą ir kėlimo elementą, nekeičiant keliamosios galios "
             "ir saugos koeficientų. Keičiame tris dalykus: rėmą, kėlimo elementą ir dėl to — skriemulį su varikliu.")


def selection_table(s, headers, rows, pick=1, bars=None):
    """Label column + 3 option columns; the chosen column is tinted, its header filled."""
    xs, ws = [80, 410, 887, 1364], [330, 477, 477, 476]
    y0 = 222
    hline(s, 80, y0, 1760)
    rect(s, px(xs[pick + 1]), px(y0), px(ws[pick + 1]), px(96), INK)
    for j, (h, tag) in enumerate(headers):
        c = PAPER if j == pick else INK
        t(s, xs[j + 1] + 22, y0 + 20, ws[j + 1] - 40, 40, h, 34, SERIF, c)
        t(s, xs[j + 1] + 22, y0 + 64, ws[j + 1] - 40, 20, tag.upper(), 16, MONO, c, spacing=0.06)
    y = y0 + 96
    hline(s, 80, y, 1760)
    for label, cells, h in rows:
        rect(s, px(xs[pick + 1]), px(y + 2), px(ws[pick + 1]), px(h - 2), TINT)
        t(s, 80, y, 320, h, label.upper(), 16, SANS, bold=True, spacing=0.14, anchor="middle")
        for j, cell in enumerate(cells):
            if isinstance(cell, tuple):   # (big number, pixel-bar fraction)
                t(s, xs[j + 1] + 22, y + 16, 300, 44, cell[0], 40, SERIF)
                on = round(20 * cell[1])
                for k in range(20):
                    rect(s, px(xs[j + 1] + 22 + k * 19), px(y + 68), px(16), px(12), INK if k < on else "#DBD8DE")
            else:
                t(s, xs[j + 1] + 22, y + 20, ws[j + 1] - 44, h - 30, cell, 24, SANS, line=1.1)
        y += h
        hline(s, 80, y, 1760, 1, INK_FAINT)
    return y


def verdict_line(s, strong, rest):
    hline(s, 80, 830, 1760)
    t(s, 80, 866, 240, 24, "SPRENDIMAS", 16, SANS, bold=True, spacing=0.18)
    v = t(s, 330, 852, 1510, 110, strong, 34, SERIF, line=1.15)
    add_run(v, rest, 27, SANS)


def s08_frame(prs, bg):
    s = new_slide(prs, bg)
    topbar(s, "Rėmo medžiagos parinkimas", "PBL 2 · Medžiagos")
    selection_table(s, [("Plienas S355", "dabar"), ("Aliuminio lydinys 6082-T6", "pasirinkta"),
                        ("Anglies pluošto kompozitas", "alternatyva")],
                    [("Tankis, kg/m³", [("7850", 1.0), ("2700", 2700 / 7850), ("1600", 1600 / 7850)], 100),
                     ("Stiprumas, MPa", ["355 (takumo riba)", "≥ 250 (takumo riba)", "≈ 1730 (išilgai pluošto)"], 66),
                     ("Tamprumo modulis, GPa", ["210", "70", "≈ 142 (išilgai pluošto)"], 66),
                     ("Savitasis stiprumas, kN·m/kg", ["≈ 45", "≈ 93", "≈ 1080"], 66),
                     ("Kaina ir gamyba", ["Pigus, lengvai suvirinamas", "Vidutinė kaina, lengvai apdirbamas, atsparus korozijai",
                                          "Brangus, sunkiai jungiamas ir remontuojamas"], 96)])
    verdict_line(s, "Aliuminio lydinys 6082-T6 — ", "beveik 3 kartus lengvesnis už plieną ir atsparus korozijai. "
                 "Tamprumo modulis 3 kartus mažesnis, todėl profiliai turės būti didesni, kad rėmas per daug nelinktų.")
    chrome(s, 8, "Šaltiniai: Weerg (2024); Wieland SMH GmbH (2018); The Engineering ToolBox (b. d.)")
    notes(s, "Rėmui palyginome plieną, aliuminį ir anglies pluoštą. Anglies pluoštas stipriausias, bet brangus ir sunkiai jungiamas. "
             "Renkamės aliuminio lydinį 6082-T6: beveik 3 kartus lengvesnis už plieną. Kadangi jis 3 kartus mažiau standus, "
             "profiliai bus didesni.")


def s09_lift(prs, bg):
    s = new_slide(prs, bg)
    topbar(s, "Kėlimo elemento parinkimas", "PBL 2 · Detalės")
    selection_table(s, [("Plieninis lynas", "dabar"), ("Dengtas plieninis diržas", "pasirinkta"),
                        ("Anglies pluošto lynas", "alternatyva")],
                    [("Sandara", ["Susukti plieniniai lynai", "Plieniniai lynai, padengti poliuretanu",
                                  "Anglies pluošto šerdis su danga"], 70),
                     ("Masė", ["Sunkūs lynai", "Diržai iki 20 % lengvesni", "500 m lifte judanti masė ≈ 27 t → 13 t"], 70),
                     ("Skriemulys", ["≥ 40 lyno skersmenų", "Apie 85 % mažesnis", "—"], 70),
                     ("Variklis", ["Didelis sukimo momentas", "Apie 70 % mažesnis", "—"], 70),
                     ("Paskirtis", ["Įprasti traukos liftai", "Liftai be mašinų patalpos (Otis Gen2, 2000)",
                                    "Labai aukšti pastatai — iki 1 km (KONE UltraRope, 2013)"], 96)])
    verdict_line(s, "Dengtas plieninis diržas — ", "sumažina ne tik lyno masę, bet ir skriemulį bei variklį. "
                 "Anglies pluošto lyno privalumas didžiausias tik labai aukštuose pastatuose.")
    chrome(s, 9, "Šaltiniai: FMLink (2005, 2013); Muñoz (2020); KONE Corporation (b. d.)")
    notes(s, "Kėlimo elementui palyginome plieninį lyną, dengtą plieninį diržą ir anglies pluošto lyną. Diržas iki 20 % lengvesnis, "
             "leidžia apie 85 % mažesnį skriemulį ir apie 70 % mažesnį variklį. Anglies pluoštas naudingiausias tik labai aukštuose pastatuose, "
             "todėl renkamės diržą.")


def s10_conclusions(prs, bg):
    s = new_slide(prs, bg)
    topbar(s, "Išvados", "TA1 · Apibendrinimas")
    items = [("RAIDA", "Keltuvų raidą lėmė dvi pagrindinės problemos: saugumas ir judančių dalių masė."),
             ("VERTINIMAS", "Saugumą iš esmės užtikrino automatinis saugos stabdys ir trinties pavara su atsvaru."),
             ("PROBLEMA", "Naujausia raidos kryptis rodo, kad didžiausias rezervas — lengvesnės judančios dalys."),
             ("SPRENDIMAS", "Siūlome rėmą gaminti iš aliuminio lydinio, o plieninį lyną pakeisti dengtu plieniniu diržu.")]
    for i, (tag, p) in enumerate(items):
        x = 80 + (i % 2) * 916
        y = 230 + (i // 2) * 330
        t(s, x, y + 50, 110, 110, str(i + 1), 96, SERIF)
        t(s, x + 110, y + 50, 700, 20, tag, 15, SANS, bold=True, spacing=0.18)
        t(s, x + 110, y + 82, 730, 190, p, 44, SERIF, line=1.05)
        hline(s, x, y + 310, 844, 1, INK_FAINT)
    chrome(s, 10)
    notes(s, "Išvados: raidą lėmė saugumas ir masė; saugumą išsprendė saugos stabdys ir trauka su atsvaru; "
             "dabar didžiausias rezervas — lengvesnės dalys; todėl siūlome aliuminio rėmą ir dengtą plieninį diržą. "
             "TA2 dalyje suprojektuosime sprendimą ir įvertinsime atitiktį standartams.")


def s11_thanks(prs, bg):
    s = new_slide(prs, bg)
    glitch(s, 80, 92, 850, 18, from_left=True)
    t(s, 1040, 330, 800, 24, "KELTUVAI · TA1", 18, SANS, bold=True, spacing=0.18, align="right")
    h = t(s, 840, 390, 1000, 400, "Ačiū", 180, SERIF, align="right", spacing=-0.02, line=0.95)
    add_run(h, "už dėmesį", 180, SERIF, new_par=True, line=0.95)
    t(s, 1240, 800, 600, 60, "Klausimai?", 44, SERIF, INK_SOFT, align="right", italic=True)
    chrome(s, 11)
    notes(s, "Ačiū už dėmesį. Mielai atsakysime į klausimus.")


REFS = [
    [("Cloux, J.-N., Pougny, J.-P., ir Menard, J.-P. (1999). ", "Elevator with reduced counterweight", " (JAV patentas Nr. 5984052). Otis Elevator Co. https://patents.google.com/patent/US5984052"),
     ("de Almeida, A., Dütschke, E., Patrão, C., Hirzel, S., ir Fong, J. (2010). Elevators and escalators: Energy performance and strategies to promote energy efficiency. In ", "Proceedings of the 6th International Conference on Improving Energy Efficiency in Commercial Buildings (IEECB Focus 2010)", ". Fraunhofer. https://publica.fraunhofer.de/handle/publica/367371"),
     ("The Elevator Museum. (b. d.). ", "Friedrich Koepe", " [Ištrauka iš J. Simmen ir J. Imorde knygos „A Cultural History of Vertical Transport“]. https://theelevatormuseum.org/koepe.php"),
     ("The Engineering ToolBox. (b. d.). ", "Engineering materials – properties", ". https://engineeringtoolbox.com/engineering-materials-properties-d_1225.html"),
     ("ERIH. (b. d.). ", "Sir William George Armstrong (1810–1900)", ". European Route of Industrial Heritage. https://www.erih.net/how-it-started/stories-about-people-biographies/biography/armstrong"),
     ("ETHW. (2015). ", "The electric elevator", ". Engineering and Technology History Wiki. https://ethw.org/The_Electric_Elevator")],
    [("FMLink. (2005, birželio 10). ", "Otis introduces Gen2 Comfort flat-belt elevator system", ". https://www.fmlink.com/otis-introduces-gen2-comfort-flat-belt-elevator-system/"),
     ("FMLink. (2013, birželio 24). ", "Don't look down: KONE unveils high-rise elevator technology that reaches one kilometer", ". https://www.fmlink.com/dont-look-down-kone-unveils-high-rise-elevator-technology-that-reaches-one-kilometer/"),
     ("Kalliomaki, J. (2022). 1927 – The year that set the direction of traction lift engineering for a century. In ", "Proceedings of the 13th Symposium on Lift & Escalator Technologies", " (T. 13, p. 141–152). https://liftescalatorlibrary.org/paper_indexing/abstract_pages/00000501.html"),
     ("Klein, C. (2015, birželio 9). ", "1,500 years later, killer animal elevator returns to Colosseum", ". History. https://www.history.com/articles/1500-years-later-killer-animal-elevator-returns-to-colosseum"),
     ("KONE Corporation. (b. d.). ", "It's all about a rope!", " https://www.kone.co.uk/stories-and-references/stories/its-all-bout-a-rope.aspx"),
     ("KONE Corporation. (2026, kovo 17). ", "A space-saving innovation that keeps getting better", ". https://www.kone.com/global/en/newsroom/stories/a-space-saving-innovation-that-keeps-getting-better.html")],
    [("Krakowski, T., ir Ruta, H. (2018). Analysis and assessment of energy efficiency of passenger lifts. ", "Advances in Science and Technology Research Journal, 12", "(3), 257–265. https://doi.org/10.12913/22998624/95165"),
     ("Muñoz, R. (2020, gruodis). 20 years of innovation. ", "Elevator World", ". https://elevatorworld.com/?p=14006"),
     ("Otis, E. G. (1861). ", "Improvement in hoisting apparatus", " (JAV patentas Nr. 31128). https://patents.google.com/patent/US31128"),
     ("Rotary Solutions. (b. d.). ", "About", ". https://rotarysolutions.com/about/"),
     ("TK Elevator. (b. d.). ", "Counterweights", " [Classroom on Demand]. https://www.tkelevator.com/us-en/tools/classroom-on-demand/counterweights.html"),
     ("Weerg. (2024). ", "Technical datasheet: Steel S355J2", ". https://weerg.com/hubfs/Datasheets/Datasheets%202024/ENG/EN_AcciaioS355J2.pdf"),
     ("Wieland SMH GmbH. (2018). ", "Material data sheet EN AW 6082", ". https://www.wieland.com/da/content/download/17142/file/EN-AW-6082_EN.pdf")],
]


def s_refs(prs, bg, page):
    s = new_slide(prs, bg)
    topbar(s, "Literatūros sąrašas", f"APA 7 · {page + 1} / 3")
    refs = REFS[page]
    half = (len(refs) + 1) // 2
    for col, chunk in enumerate((refs[:half], refs[half:])):
        x = 80 + col * 916
        box = None
        for pre, ital, post in chunk:
            if box is None:
                box = t(s, x, 232, 844, 740, pre, 23, SANS, line=1.12)
            else:
                add_run(box, pre, 23, SANS, new_par=True, line=1.12, space_before=10)
            add_run(box, ital, 23, SANS, italic=True)
            add_run(box, post, 23, SANS)
    chrome(s, 12 + page)


def main():
    theme = resolved_theme()
    bg = grid_background()
    prs = Presentation(theme=theme, strict_theme=True)
    for build in (s01_cover, s02_why, s03_history, s04_schemes, s05_verdict, s06_problem, s07_task,
                  s08_frame, s09_lift, s10_conclusions, s11_thanks):
        build(prs, bg)
    for page in range(3):
        s_refs(prs, bg, page)
    prs.save(OUT)
    print("saved", OUT, "slides:", len(prs.slides))


if __name__ == "__main__":
    main()
